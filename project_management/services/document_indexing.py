import hashlib
import logging
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from difflib import SequenceMatcher
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Iterable, List, Optional

from django.conf import settings
from django.contrib.postgres.search import (
    SearchQuery,
    SearchRank,
    SearchVector,
    TrigramSimilarity,
    TrigramWordSimilarity,
)
from django.core.files.storage import default_storage
from django.db import IntegrityError, models, transaction
from django.db.models import Case, F, OuterRef, Q, Subquery, Value, When
from django.db.models.functions import Coalesce, Greatest
from django.utils import timezone

from project_management.models import (
    Document,
    DocumentContentSegment,
    DocumentIndexAttempt,
    DocumentIndexJob,
    DocumentSearchIndex,
)
from project_management.filters import DocumentFilter
from project_management.services.document_access import get_accessible_document_queryset

logger = logging.getLogger(__name__)


class DocumentIndexingError(Exception):
    def __init__(self, message: str, category: str = "extraction_error"):
        super().__init__(message)
        self.category = category


@dataclass
class ExtractedSegment:
    segment_type: str
    sequence: int
    location_label: str
    text: str
    attachment_name: str = ""
    ocr_confidence: Optional[float] = None


SUPPORTED_TEXT_EXTENSIONS = {
    ".txt",
    ".csv",
    ".md",
    ".json",
    ".xml",
    ".html",
    ".htm",
}
SUPPORTED_TIKA_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".rtf",
    ".odt",
    ".xls",
    ".xlsx",
    ".ods",
    ".ppt",
    ".pptx",
    ".odp",
}
SUPPORTED_IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".webp",
}
SUPPORTED_EMAIL_EXTENSIONS = {".eml", ".msg"}
STORED_ONLY_EXTENSIONS = {".zip"}
SUPPORTED_CAD_INTERCHANGE_EXTENSIONS = {".dxf", ".step", ".stp", ".iges", ".igs"}
SUPPORTED_PROPRIETARY_CAD_EXTENSIONS = {
    ".asm",
    ".catdrawing",
    ".dwg",
    ".dwt",
    ".dws",
    ".catpart",
    ".catproduct",
    ".cgr",
    ".drw",
    ".frm",
    ".iam",
    ".idw",
    ".ipn",
    ".ipt",
    ".mfg",
    ".model",
    ".prt",
    ".rfa",
    ".rte",
    ".rvt",
    ".sec",
    ".neu",
}

MIME_TYPE_EXTENSIONS = {
    "application/json": ".json",
    "application/msword": ".doc",
    "application/pdf": ".pdf",
    "application/rtf": ".rtf",
    "application/vnd.ms-excel": ".xls",
    "application/vnd.ms-powerpoint": ".ppt",
    "application/vnd.oasis.opendocument.presentation": ".odp",
    "application/vnd.oasis.opendocument.spreadsheet": ".ods",
    "application/vnd.oasis.opendocument.text": ".odt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/xml": ".xml",
    "application/x-zip-compressed": ".zip",
    "application/zip": ".zip",
    "image/bmp": ".bmp",
    "image/gif": ".gif",
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/tiff": ".tiff",
    "image/webp": ".webp",
    "message/rfc822": ".eml",
    "text/csv": ".csv",
    "text/html": ".html",
    "text/markdown": ".md",
    "text/plain": ".txt",
    "text/rtf": ".rtf",
    "text/xml": ".xml",
}

SUPPORTED_DOCUMENT_EXTENSIONS = (
    SUPPORTED_TEXT_EXTENSIONS
    | SUPPORTED_TIKA_EXTENSIONS
    | SUPPORTED_IMAGE_EXTENSIONS
    | SUPPORTED_EMAIL_EXTENSIONS
    | STORED_ONLY_EXTENSIONS
    | SUPPORTED_CAD_INTERCHANGE_EXTENSIONS
    | SUPPORTED_PROPRIETARY_CAD_EXTENSIONS
)

# PostgreSQL limits a single tsvector value to roughly 1 MB. Keep source
# segments comfortably below that boundary before generating search vectors.
SEARCH_VECTOR_SEGMENT_MAX_BYTES = 200_000


def document_indexing_enabled() -> bool:
    return bool(getattr(settings, "DOCUMENT_INDEXING_ENABLED", False))


def document_ocr_enabled() -> bool:
    return bool(getattr(settings, "DOCUMENT_OCR_ENABLED", False))


def enqueue_document_indexing(document_id: int, *, force: bool = False) -> Optional[DocumentIndexJob]:
    if not document_indexing_enabled():
        return None

    try:
        document = Document.objects.select_related("folder").get(id=document_id)
    except Document.DoesNotExist:
        return None

    if not document.file:
        return None

    content_hash = compute_document_hash(document)
    index, _ = DocumentSearchIndex.objects.get_or_create(document=document)

    content_unchanged = index.content_hash == content_hash
    parser = getattr(
        settings,
        "DOCUMENT_INDEXING_PARSER_VERSION",
        "document-search-v2",
    )
    parser_current = index.parser_version == parser
    if (
        not force
        and content_unchanged
        and parser_current
        and index.status in [DocumentSearchIndex.Status.SEARCHABLE, DocumentSearchIndex.Status.UNSUPPORTED]
    ):
        return None

    active_job = DocumentIndexJob.objects.filter(
        document=document,
        content_hash=content_hash,
        status__in=[
            DocumentIndexJob.Status.QUEUED,
            DocumentIndexJob.Status.RUNNING,
            DocumentIndexJob.Status.RETRYING,
        ],
    ).first()
    if active_job and not force:
        return active_job

    max_attempts = int(getattr(settings, "DOCUMENT_INDEXING_MAX_ATTEMPTS", 3))
    if active_job:
        active_job.status = DocumentIndexJob.Status.CANCELLED
        active_job.completed_at = timezone.now()
        active_job.save(update_fields=["status", "completed_at", "updated_at"])

    try:
        # Keep a duplicate-enqueue race from marking a caller's outer transaction
        # as broken. The database constraint remains the source of truth.
        with transaction.atomic():
            job = DocumentIndexJob.objects.create(
                document=document,
                content_hash=content_hash,
                max_attempts=max_attempts,
            )
    except IntegrityError:
        return DocumentIndexJob.objects.filter(
            document=document,
            content_hash=content_hash,
            status__in=[
                DocumentIndexJob.Status.QUEUED,
                DocumentIndexJob.Status.RUNNING,
                DocumentIndexJob.Status.RETRYING,
            ],
        ).first()

    if not content_unchanged:
        DocumentContentSegment.objects.filter(document=document).delete()
    index.status = DocumentSearchIndex.Status.PENDING
    index.content_hash = content_hash
    index.error = ""
    index.failure_category = ""
    index.save(update_fields=["status", "content_hash", "error", "failure_category", "updated_at"])

    transaction.on_commit(lambda: dispatch_document_indexing_job(job.id))
    return job


def dispatch_document_indexing_job(job_id: int) -> None:
    try:
        from project_management.tasks import process_document_index_job

        job = DocumentIndexJob.objects.filter(id=job_id).only("available_at").first()
        if not job:
            return
        process_document_index_job.apply_async(
            args=[job_id],
            queue=getattr(settings, "DOCUMENT_INDEXING_QUEUE", "document_indexing"),
            eta=job.available_at if job.available_at > timezone.now() else None,
        )
    except Exception:
        logger.exception("Failed to dispatch document indexing job %s", job_id)


def compute_document_hash(document: Document) -> str:
    sha = hashlib.sha256()
    file_obj = document.file
    try:
        file_obj.open("rb")
        for chunk in iter(lambda: file_obj.read(1024 * 1024), b""):
            sha.update(chunk)
    finally:
        try:
            file_obj.close()
        except Exception:
            pass
    return sha.hexdigest()


def process_index_job(job_id: int) -> None:
    started = time.monotonic()
    with transaction.atomic():
        job = (
            DocumentIndexJob.objects.select_for_update()
            .select_related("document", "document__folder")
            .filter(id=job_id)
            .first()
        )
        if not job or job.status in [
            DocumentIndexJob.Status.COMPLETED,
            DocumentIndexJob.Status.CANCELLED,
        ]:
            return
        if job.status == DocumentIndexJob.Status.RUNNING:
            return
        if job.available_at > timezone.now():
            transaction.on_commit(lambda: dispatch_document_indexing_job(job_id))
            return

        document = job.document
        if compute_document_hash(document) != job.content_hash:
            job.status = DocumentIndexJob.Status.CANCELLED
            job.completed_at = timezone.now()
            job.save(update_fields=["status", "completed_at", "updated_at"])
            return

        job.status = DocumentIndexJob.Status.RUNNING
        job.started_at = timezone.now()
        job.attempt_count = F("attempt_count") + 1
        job.save(update_fields=["status", "started_at", "attempt_count", "updated_at"])

    document = job.document
    job.refresh_from_db()
    index, _ = DocumentSearchIndex.objects.get_or_create(document=document)
    source_size = getattr(document.file, "size", 0) or 0
    max_bytes = int(getattr(settings, "DOCUMENT_INDEXING_MAX_FILE_SIZE_MB", 100)) * 1024 * 1024

    index.status = DocumentSearchIndex.Status.PROCESSING
    index.attempt_count = job.attempt_count
    index.save(update_fields=["status", "attempt_count", "updated_at"])

    extracted_size = 0
    parser = getattr(settings, "DOCUMENT_INDEXING_PARSER_VERSION", "document-search-v2")
    try:
        if source_size > max_bytes:
            raise DocumentIndexingError("File is larger than the configured indexing limit.", "oversized")

        segments, metadata = extract_document(document)
        extracted_size = sum(len(segment.text.encode("utf-8")) for segment in segments)
        save_extracted_segments(index, document, segments)

        if metadata.get("stored_only"):
            index.status = DocumentSearchIndex.Status.UNSUPPORTED
        elif metadata.get("metadata_only"):
            index.status = DocumentSearchIndex.Status.PARTIAL
        elif not segments:
            index.status = DocumentSearchIndex.Status.UNSUPPORTED
        elif metadata.get("partial"):
            index.status = DocumentSearchIndex.Status.PARTIAL
        else:
            index.status = DocumentSearchIndex.Status.SEARCHABLE
        index.mime_type = metadata.get("mime_type", "")
        index.language = metadata.get("language", "")
        index.extraction_method = metadata.get("extraction_method", DocumentSearchIndex.ExtractionMethod.NONE)
        index.page_count = metadata.get("page_count", len(segments))
        index.parser_version = parser
        index.error = ""
        index.failure_category = ""
        index.indexed_at = timezone.now()
        index.save()

        job.status = DocumentIndexJob.Status.COMPLETED
        job.completed_at = timezone.now()
        job.error = ""
        job.failure_category = ""
        job.save(update_fields=["status", "completed_at", "error", "failure_category", "updated_at"])
        record_attempt(job, document, "completed", parser, started, source_size, extracted_size)
    except DocumentIndexingError as exc:
        handle_index_failure(job, index, document, exc, parser, started, source_size, extracted_size)
    except Exception as exc:
        handle_index_failure(
            job,
            index,
            document,
            DocumentIndexingError(str(exc), "unexpected_error"),
            parser,
            started,
            source_size,
            extracted_size,
        )


def handle_index_failure(job, index, document, exc, parser, started, source_size, extracted_size):
    should_retry = job.attempt_count < job.max_attempts and exc.category not in {
        "unsupported",
        "oversized",
        "encrypted",
        "page_limit",
    }
    job.status = DocumentIndexJob.Status.RETRYING if should_retry else DocumentIndexJob.Status.FAILED
    job.error = str(exc)[:2000]
    job.failure_category = exc.category
    job.available_at = timezone.now() + timezone.timedelta(minutes=min(30, 2 ** job.attempt_count))
    job.save(update_fields=["status", "error", "failure_category", "available_at", "updated_at"])

    index.status = DocumentSearchIndex.Status.FAILED if not should_retry else DocumentSearchIndex.Status.PENDING
    index.error = str(exc)[:2000]
    index.failure_category = exc.category
    index.save(update_fields=["status", "error", "failure_category", "updated_at"])
    record_attempt(job, document, "failed", parser, started, source_size, extracted_size, exc.category, str(exc))

    if should_retry:
        dispatch_document_indexing_job(job.id)


def record_attempt(
    job,
    document,
    status,
    parser,
    started,
    source_size,
    extracted_size,
    failure_category="",
    error="",
):
    DocumentIndexAttempt.objects.create(
        job=job,
        document=document,
        status=status,
        parser=parser,
        failure_category=failure_category,
        error=error[:2000],
        duration_ms=max(0, int((time.monotonic() - started) * 1000)),
        source_size_bytes=source_size,
        extracted_size_bytes=extracted_size,
    )


def save_extracted_segments(index, document, segments: Iterable[ExtractedSegment]) -> None:
    DocumentContentSegment.objects.filter(document=document).delete()
    normalized_segments = []
    metadata_text = normalize_extracted_text(
        " ".join(
            str(value)
            for value in [
                document.file_name,
                document.full_document_number,
                document.identification_number,
                document.document_type,
                document.status,
                document.category,
                document.discipline,
                document.folder.name,
            ]
            if value
        )
    )
    if metadata_text:
        normalized_segments.append(
            ExtractedSegment("document", 0, "Metadata", metadata_text)
        )
    for segment in segments:
        normalized_text = normalize_extracted_lines(segment.text)
        if normalized_text:
            chunks = split_search_vector_text(normalized_text)
            for chunk_number, chunk in enumerate(chunks, start=1):
                location_label = segment.location_label
                if len(chunks) > 1:
                    location_label = (
                        f"{segment.location_label} ({chunk_number}/{len(chunks)})"
                    )[:255]
                normalized_segments.append(
                    ExtractedSegment(
                        segment.segment_type,
                        segment.sequence,
                        location_label,
                        chunk,
                        segment.attachment_name,
                        segment.ocr_confidence,
                    )
                )

    segment_objects = [
        DocumentContentSegment(
            document=document,
            index=index,
            segment_type=segment.segment_type,
            sequence=segment.sequence,
            location_label=segment.location_label,
            text=segment.text,
            attachment_name=segment.attachment_name,
            ocr_confidence=segment.ocr_confidence,
        )
        for segment in normalized_segments
    ]
    DocumentContentSegment.objects.bulk_create(segment_objects, batch_size=100)
    DocumentContentSegment.objects.filter(document=document).update(
        search_vector=SearchVector("text", config="simple", weight="C")
    )


def extract_file_content(file_field, display_name: str = ""):
    storage_name = os.path.basename(file_field.name)
    name = display_name or storage_name
    mime_type, _ = mimetypes.guess_type(name)

    with tempfile.TemporaryDirectory(prefix="doc-index-") as tmp_dir:
        local_path = Path(tmp_dir) / safe_tmp_filename(name)
        with default_storage.open(file_field.name, "rb") as source, open(local_path, "wb") as target:
            shutil.copyfileobj(source, target, length=1024 * 1024)
        mime_type = detect_mime_type(local_path, mime_type)
        extension = resolve_extraction_extension(name, storage_name, mime_type)
        local_path = ensure_extraction_path_suffix(local_path, extension)

        if extension in SUPPORTED_TEXT_EXTENSIONS:
            text = read_text_file(local_path)
            if not normalize_extracted_text(text):
                raise DocumentIndexingError("No extractable text found.", "empty")
            return [ExtractedSegment("document", 1, "Document", text)], {
                "mime_type": mime_type or "text/plain",
                "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
                "page_count": 1,
            }

        if extension in SUPPORTED_EMAIL_EXTENSIONS:
            return extract_email(local_path, extension, mime_type)

        if extension in STORED_ONLY_EXTENSIONS:
            return [], {
                "mime_type": mime_type or "application/zip",
                "extraction_method": DocumentSearchIndex.ExtractionMethod.NONE,
                "page_count": 0,
                "stored_only": True,
            }

        if extension == ".pdf":
            return extract_pdf(local_path, mime_type)

        if extension == ".docx":
            return extract_docx(local_path, mime_type)

        if extension == ".xlsx":
            return extract_xlsx(local_path, mime_type)

        if extension == ".pptx":
            return extract_pptx(local_path, mime_type)

        if extension in SUPPORTED_TIKA_EXTENSIONS:
            text = run_tika(local_path)
            if text.strip():
                return [ExtractedSegment("document", 1, "Document", text)], {
                    "mime_type": mime_type or "",
                    "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
                    "page_count": estimate_page_count(text),
                }
            raise DocumentIndexingError("No extractable text found.", "empty")

        if extension in SUPPORTED_CAD_INTERCHANGE_EXTENSIONS:
            return extract_cad_interchange(local_path, extension, mime_type)

        if extension in SUPPORTED_PROPRIETARY_CAD_EXTENSIONS:
            return [], cad_metadata_only(extension, mime_type)

        if extension in SUPPORTED_IMAGE_EXTENSIONS and document_ocr_enabled():
            return extract_image_with_ocr(local_path, mime_type)

        raise DocumentIndexingError(f"Unsupported document format: {extension or 'unknown'}", "unsupported")


def extract_document(document: Document):
    storage_name = os.path.basename(document.file.name)
    return extract_file_content(document.file, document.file_name or storage_name)


def extract_pdf(path: Path, mime_type: Optional[str]):
    native_segments, total_pages = read_pdf_segments(path)
    has_missing_text = len(native_segments) < total_pages
    if not has_missing_text:
        return native_segments, {
            "mime_type": mime_type or "application/pdf",
            "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
            "page_count": total_pages,
        }

    if not document_ocr_enabled():
        if native_segments:
            return native_segments, {
                "mime_type": mime_type or "application/pdf",
                "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
                "page_count": total_pages,
                "partial": True,
            }
        raise DocumentIndexingError("No extractable PDF text found and OCR is disabled.", "empty")

    ocrmypdf_bin = shutil.which("ocrmypdf")
    if not ocrmypdf_bin:
        if native_segments:
            return native_segments, {
                "mime_type": mime_type or "application/pdf",
                "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
                "page_count": total_pages,
                "partial": True,
            }
        raise DocumentIndexingError("OCRmyPDF is not installed.", "missing_dependency")

    output_path = path.with_name(f"{path.stem}-ocr.pdf")
    run_command(
        [
            ocrmypdf_bin,
            "--skip-text",
            "--rotate-pages",
            "--deskew",
            "--language",
            "eng+deu+fra",
            str(path),
            str(output_path),
        ]
    )
    ocr_segments, ocr_page_count = read_pdf_segments(output_path)
    if not ocr_segments:
        raise DocumentIndexingError("OCR did not produce searchable PDF text.", "ocr_failed")
    return ocr_segments, {
        "mime_type": mime_type or "application/pdf",
        "language": "eng+deu+fra",
        "extraction_method": (
            DocumentSearchIndex.ExtractionMethod.MIXED
            if native_segments
            else DocumentSearchIndex.ExtractionMethod.OCR
        ),
        "page_count": ocr_page_count,
    }


def read_pdf_segments(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentIndexingError("pypdf is not installed.", "missing_dependency") from exc

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentIndexingError("Encrypted PDF cannot be indexed.", "encrypted")

    page_limit = int(getattr(settings, "DOCUMENT_INDEXING_MAX_PAGES", 1000))
    if len(reader.pages) > page_limit:
        raise DocumentIndexingError("PDF exceeds the configured page limit.", "page_limit")

    segments = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            extracted_text = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            extracted_text = page.extract_text() or ""
        text = normalize_extracted_lines(extracted_text)
        if text:
            segments.append(ExtractedSegment("page", page_number, f"Page {page_number}", text))
    return segments, len(reader.pages)


def extract_docx(path: Path, mime_type: Optional[str]):
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise DocumentIndexingError("python-docx is not installed.", "missing_dependency") from exc

    document = DocxDocument(str(path))
    section_text = []
    seen_parts = set()
    for section in document.sections:
        for part in (section.header, section.footer):
            part_key = str(part.part.partname)
            if part_key in seen_parts:
                continue
            seen_parts.add(part_key)
            section_text.extend(paragraph.text for paragraph in part.paragraphs)
            section_text.extend(table_text(part.tables))

    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + table_text(document.tables)
        + section_text
    )
    if not normalize_extracted_text(text):
        raise DocumentIndexingError("No extractable Word text found.", "empty")
    return [ExtractedSegment("document", 1, "Document", text)], {
        "mime_type": mime_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
        "page_count": 1,
    }


def extract_xlsx(path: Path, mime_type: Optional[str]):
    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentIndexingError("openpyxl is not installed.", "missing_dependency") from exc

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentIndexingError(
            "Spreadsheet file could not be opened.",
            "invalid_file",
        ) from exc
    segments = []
    for sequence, sheet in enumerate(workbook.worksheets, start=1):
        rows = [
            " ".join(str(cell) for cell in row if cell is not None)
            for row in sheet.iter_rows(values_only=True)
        ]
        text = "\n".join(row for row in rows if row.strip())
        if normalize_extracted_text(text):
            segments.append(ExtractedSegment("sheet", sequence, sheet.title, text))
    if not segments:
        raise DocumentIndexingError("No extractable spreadsheet text found.", "empty")
    return segments, {
        "mime_type": mime_type or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
        "page_count": len(workbook.worksheets),
    }


def extract_pptx(path: Path, mime_type: Optional[str]):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentIndexingError("python-pptx is not installed.", "missing_dependency") from exc

    presentation = Presentation(str(path))
    segments = []
    for sequence, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and normalize_extracted_text(shape.text)
        ]
        notes = getattr(slide, "notes_slide", None)
        if notes and getattr(notes, "notes_text_frame", None):
            notes_text = notes.notes_text_frame.text
            if normalize_extracted_text(notes_text):
                text_parts.append(notes_text)
        text = "\n".join(text_parts)
        if normalize_extracted_text(text):
            segments.append(ExtractedSegment("slide", sequence, f"Slide {sequence}", text))
    if not segments:
        raise DocumentIndexingError("No extractable presentation text found.", "empty")
    return segments, {
        "mime_type": mime_type or "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
        "page_count": len(presentation.slides),
    }


def run_tika(path: Path) -> str:
    tika_bin = shutil.which("tika")
    if tika_bin:
        return run_command([tika_bin, "--text", str(path)])

    java_bin = shutil.which("java")
    tika_jar = os.environ.get("TIKA_JAR")
    if java_bin and tika_jar and os.path.exists(tika_jar):
        return run_command([java_bin, "-jar", tika_jar, "--text", str(path)])

    return run_python_fallback(path)


def run_python_fallback(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".docx":
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            return "\n".join([paragraph.text for paragraph in doc.paragraphs] + table_text(doc.tables))
        except ImportError as exc:
            raise DocumentIndexingError("python-docx is not installed.", "missing_dependency") from exc
    if extension in {".xlsx", ".xls"}:
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    rows.append(" ".join(str(cell) for cell in row if cell is not None))
            return "\n".join(rows)
        except ImportError as exc:
            raise DocumentIndexingError("openpyxl is not installed.", "missing_dependency") from exc
    if extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise DocumentIndexingError("Encrypted PDF cannot be indexed.", "encrypted")
            texts = []
            for idx, page in enumerate(reader.pages[: int(getattr(settings, "DOCUMENT_INDEXING_MAX_PAGES", 1000))]):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    texts.append(f"Page {idx + 1}\n{page_text}")
            return "\n\n".join(texts)
        except ImportError as exc:
            raise DocumentIndexingError("pypdf is not installed.", "missing_dependency") from exc
    raise DocumentIndexingError("Native extractor dependency is not installed for this file.", "missing_dependency")


def extract_image_with_ocr(path: Path, mime_type: Optional[str]):
    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        raise DocumentIndexingError("Tesseract is not installed.", "missing_dependency")
    ocr_source = prepare_image_for_ocr(path)
    output_base = path.with_suffix("")
    language = "eng+deu+fra"
    run_command(
        [tesseract_bin, str(ocr_source), str(output_base), "-l", language],
        text_output=False,
    )
    txt_path = output_base.with_suffix(".txt")
    if not txt_path.exists():
        raise DocumentIndexingError("OCR did not produce text output.", "ocr_failed")
    text = read_text_file(txt_path)
    if not normalize_extracted_text(text):
        raise DocumentIndexingError("OCR did not produce searchable text.", "ocr_failed")
    return [ExtractedSegment("document", 1, "OCR text", text)], {
        "mime_type": mime_type or "",
        "language": language,
        "extraction_method": DocumentSearchIndex.ExtractionMethod.OCR,
        "page_count": 1,
    }


def prepare_image_for_ocr(path: Path) -> Path:
    if path.suffix.lower() != ".gif":
        return path

    try:
        from PIL import Image
    except ImportError as exc:
        raise DocumentIndexingError(
            "Pillow is required to process GIF images.",
            "missing_dependency",
        ) from exc

    output_path = path.with_name(f"{path.stem}-first-frame.png")
    try:
        with Image.open(path) as image:
            image.seek(0)
            image.convert("RGB").save(output_path, format="PNG")
    except Exception as exc:
        raise DocumentIndexingError(
            "GIF could not be converted for OCR.",
            "ocr_failed",
        ) from exc
    return output_path


def extract_cad_interchange(
    path: Path,
    extension: str,
    mime_type: Optional[str],
):
    raw = path.read_bytes()
    if not raw or raw.count(b"\x00") > max(8, len(raw) // 100):
        return [], cad_metadata_only(extension, mime_type)

    text = raw.decode("utf-8", errors="ignore")
    if not normalize_extracted_text(text):
        text = raw.decode("latin-1", errors="ignore")

    if extension == ".dxf":
        values = extract_dxf_text_values(text)
    elif extension in {".step", ".stp"}:
        values = extract_step_text_values(text)
    else:
        values = extract_iges_text_values(text)

    searchable_text = normalize_extracted_text("\n".join(values))
    if not searchable_text:
        return [], cad_metadata_only(extension, mime_type)

    return [
        ExtractedSegment("document", 1, "CAD properties", searchable_text)
    ], {
        "mime_type": mime_type or cad_mime_type(extension),
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
        "page_count": 1,
    }


def cad_metadata_only(extension: str, mime_type: Optional[str]) -> dict:
    return {
        "mime_type": mime_type or cad_mime_type(extension),
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NONE,
        "page_count": 0,
        "partial": True,
        "metadata_only": True,
    }


def cad_mime_type(extension: str) -> str:
    return {
        ".dxf": "image/vnd.dxf",
        ".step": "model/step",
        ".stp": "model/step",
        ".iges": "model/iges",
        ".igs": "model/iges",
    }.get(extension, "application/octet-stream")


def extract_dxf_text_values(text: str) -> List[str]:
    lines = [line.strip() for line in text.splitlines()]
    searchable_codes = {1, 2, 3, 6, 7, 8, 100, 102, *range(300, 310)}
    values = []
    for index in range(0, len(lines) - 1, 2):
        try:
            group_code = int(lines[index])
        except ValueError:
            continue
        value = lines[index + 1]
        if group_code in searchable_codes and is_meaningful_cad_value(value):
            values.append(value)
    return unique_values(values)


def extract_step_text_values(text: str) -> List[str]:
    values = [
        match.replace("''", "'")
        for match in re.findall(r"'((?:''|[^'])+)'", text)
        if is_meaningful_cad_value(match)
    ]
    return unique_values(values)


def extract_iges_text_values(text: str) -> List[str]:
    values = []
    for line in text.splitlines():
        # IGES Start-section records contain human-readable file descriptions
        # in columns 1-72, followed by the section marker and sequence number.
        if len(line) >= 73 and line[72].upper() == "S":
            value = line[:72].strip()
            if is_meaningful_cad_value(value):
                values.append(value)

    for match in re.finditer(r"(\d{1,6})H", text, flags=re.IGNORECASE):
        length = int(match.group(1))
        start = match.end()
        value = text[start : start + length].strip()
        if is_meaningful_cad_value(value):
            values.append(value)
    return unique_values(values)


def is_meaningful_cad_value(value: str) -> bool:
    normalized = normalize_extracted_text(value)
    if len(normalized) < 2 or normalized in {"$", "*"}:
        return False
    return any(character.isalpha() for character in normalized)


def unique_values(values: Iterable[str]) -> List[str]:
    seen = set()
    unique = []
    for value in values:
        normalized = normalize_extracted_text(value)
        key = normalized.casefold()
        if normalized and key not in seen:
            seen.add(key)
            unique.append(normalized)
    return unique


def extract_email(path: Path, extension: str, mime_type: Optional[str]):
    if extension == ".msg":
        text = run_tika(path)
        if not text.strip():
            raise DocumentIndexingError("MSG text could not be extracted.", "empty")
        return [ExtractedSegment("email_body", 1, "Email body", text)], {
            "mime_type": mime_type or "application/vnd.ms-outlook",
            "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
            "page_count": 1,
        }

    with open(path, "rb") as source:
        message = BytesParser(policy=policy.default).parse(source)
    parts = []
    sequence = 1
    if message.get_body(preferencelist=("plain",)):
        parts.append(
            ExtractedSegment(
                "email_body",
                sequence,
                "Email body",
                message.get_body(preferencelist=("plain",)).get_content(),
            )
        )
        sequence += 1
    attachment_limit = int(getattr(settings, "DOCUMENT_INDEXING_MAX_ATTACHMENTS", 50))
    for attachment in list(message.iter_attachments())[:attachment_limit]:
        filename = attachment.get_filename() or f"attachment-{sequence}"
        payload = attachment.get_payload(decode=True)
        if not payload:
            continue
        text = extract_email_attachment_text(path.parent, filename, payload, attachment.get_content_type())
        if text.strip():
            parts.append(ExtractedSegment("attachment", sequence, filename, text, attachment_name=filename))
            sequence += 1
    return parts, {
        "mime_type": mime_type or "message/rfc822",
        "extraction_method": DocumentSearchIndex.ExtractionMethod.NATIVE,
        "page_count": len(parts),
    }


def extract_email_attachment_text(
    temp_dir: Path,
    filename: str,
    payload: bytes,
    mime_type: Optional[str],
) -> str:
    extension = Path(filename).suffix.lower()
    attachment_path = temp_dir / f"attachment-{hashlib.sha256(payload).hexdigest()[:12]}{extension}"
    attachment_path.write_bytes(payload)

    try:
        if extension in SUPPORTED_TEXT_EXTENSIONS or (mime_type or "").startswith("text/"):
            return read_text_file(attachment_path)
        if extension == ".pdf":
            segments, _ = extract_pdf(attachment_path, mime_type)
        elif extension == ".docx":
            segments, _ = extract_docx(attachment_path, mime_type)
        elif extension == ".xlsx":
            segments, _ = extract_xlsx(attachment_path, mime_type)
        elif extension == ".pptx":
            segments, _ = extract_pptx(attachment_path, mime_type)
        elif extension in SUPPORTED_TIKA_EXTENSIONS:
            return run_tika(attachment_path)
        elif extension in SUPPORTED_CAD_INTERCHANGE_EXTENSIONS:
            segments, _ = extract_cad_interchange(
                attachment_path,
                extension,
                mime_type,
            )
        elif extension in SUPPORTED_PROPRIETARY_CAD_EXTENSIONS:
            return filename
        elif extension in SUPPORTED_IMAGE_EXTENSIONS and document_ocr_enabled():
            segments, _ = extract_image_with_ocr(attachment_path, mime_type)
        else:
            return ""
    except DocumentIndexingError:
        logger.info("Skipped unsupported email attachment %s", safe_tmp_filename(filename))
        return ""
    return "\n".join(segment.text for segment in segments)


def run_command(command: List[str], *, text_output: bool = True) -> str:
    timeout = int(getattr(settings, "DOCUMENT_INDEXING_TIMEOUT_SECONDS", 900))
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=text_output,
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentIndexingError("Extractor timed out.", "timeout") from exc
    if completed.returncode != 0:
        stderr = completed.stderr if text_output else completed.stderr.decode("utf-8", errors="ignore")
        raise DocumentIndexingError(stderr.strip() or "Extractor command failed.", "extraction_error")
    return completed.stdout if text_output else ""


def table_text(tables):
    lines = []
    for table in tables:
        for row in table.rows:
            lines.append(" ".join(cell.text for cell in row.cells))
    return lines


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def detect_mime_type(path: Path, fallback: Optional[str] = None) -> str:
    file_bin = shutil.which("file")
    if file_bin:
        try:
            detected = run_command([file_bin, "--brief", "--mime-type", str(path)]).strip()
            if detected:
                return detected
        except DocumentIndexingError:
            pass
    return fallback or "application/octet-stream"


def resolve_extraction_extension(display_name: str, storage_name: str, mime_type: str) -> str:
    for name in (display_name, storage_name):
        extension = Path(name or "").suffix.lower()
        if extension in SUPPORTED_DOCUMENT_EXTENSIONS:
            return extension

    return MIME_TYPE_EXTENSIONS.get((mime_type or "").lower(), "")


def normalize_extracted_text(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\x00", " ")).strip()


def normalize_extracted_lines(text: str) -> str:
    normalized = (
        (text or "")
        .replace("\x00", " ")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )
    lines = [
        re.sub(r"[^\S\n]+", " ", line).strip()
        for line in normalized.split("\n")
    ]
    return "\n".join(line for line in lines if line)


def split_search_vector_text(
    text: str,
    max_bytes: int = SEARCH_VECTOR_SEGMENT_MAX_BYTES,
) -> List[str]:
    """Split text at UTF-8-safe boundaries before PostgreSQL vectorization."""
    if max_bytes <= 0:
        raise ValueError("max_bytes must be greater than zero")

    encoded = (text or "").encode("utf-8")
    if len(encoded) <= max_bytes:
        return [text] if text else []

    chunks = []
    start = 0
    while start < len(encoded):
        end = min(start + max_bytes, len(encoded))
        while end < len(encoded) and end > start and encoded[end] & 0xC0 == 0x80:
            end -= 1

        candidate = encoded[start:end]
        if end < len(encoded):
            whitespace_boundary = max(
                candidate.rfind(b" "),
                candidate.rfind(b"\n"),
                candidate.rfind(b"\t"),
            )
            if whitespace_boundary >= max_bytes // 2:
                end = start + whitespace_boundary + 1

        chunk = encoded[start:end].decode("utf-8").strip()
        if chunk:
            chunks.append(chunk)
        start = end

    return chunks


def estimate_page_count(text: str) -> int:
    page_markers = re.findall(r"\bPage\s+\d+\b", text or "", flags=re.IGNORECASE)
    return max(1, len(page_markers))


def safe_tmp_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]", "_", name or "document").lstrip("._")
    return cleaned[:180] or "document"


def ensure_extraction_path_suffix(path: Path, extension: str) -> Path:
    """Give extension-sensitive parsers a path matching the detected file type."""
    if not extension or path.suffix.lower() == extension.lower():
        return path

    target = path.with_name(f"{path.name}{extension}")
    path.replace(target)
    return target


def accessible_document_queryset(request):
    project_id = request.query_params.get("project_id")
    queryset = get_accessible_document_queryset(
        request,
        project_id=project_id,
    ).select_related(
        "folder",
        "folder__project",
        "folder__company_profile",
        "author",
        "search_index",
    )
    queryset = DocumentFilter(request.query_params, queryset=queryset).qs
    scope_filters = {
        "project_id": "folder__project_id",
        "company_profile_id": "folder__company_profile_id",
        "freelancer_id": "folder__freelancer_id",
    }
    for parameter, field_name in scope_filters.items():
        value = request.query_params.get(parameter)
        if value:
            queryset = queryset.filter(**{field_name: value})
    return queryset


DOCUMENT_METADATA_SEARCH_FIELDS = (
    "file_name",
    "full_document_number",
    "identification_number",
    "document_type",
    "status",
    "category",
    "discipline",
    "folder__name",
)


def document_metadata_text(document: Document) -> str:
    values = [
        document.file_name,
        document.full_document_number,
        document.identification_number,
        document.document_type,
        document.status,
        document.category,
        document.discipline,
        document.folder.name if document.folder_id else None,
    ]
    return normalize_extracted_text(" ".join(str(value) for value in values if value))


def build_document_metadata_match(query: str) -> Q:
    phrase_match = Q()
    for field_name in DOCUMENT_METADATA_SEARCH_FIELDS:
        phrase_match |= Q(**{f"{field_name}__icontains": query})
    phrase_match |= Q(file_name__unaccent__icontains=query)

    terms = extract_search_terms(query)
    if not terms:
        return phrase_match

    token_match = Q()
    for term in terms:
        term_match = Q()
        for field_name in DOCUMENT_METADATA_SEARCH_FIELDS:
            term_match |= Q(**{f"{field_name}__icontains": term})
        term_match |= Q(file_name__unaccent__icontains=term)
        token_match &= term_match
    return phrase_match | token_match


def build_document_metadata_rank(query: str):
    metadata_match = build_document_metadata_match(query)
    return Case(
        When(file_name__iexact=query, then=Value(1.0)),
        When(file_name__icontains=query, then=Value(0.95)),
        When(full_document_number__icontains=query, then=Value(0.9)),
        When(identification_number__icontains=query, then=Value(0.85)),
        When(metadata_match, then=Value(0.75)),
        default=Value(0.0),
        output_field=models.FloatField(),
    )


def _ordered_document_results(documents, ordering: str):
    ordering_fields = {
        "relevance": ("-rank", "id"),
        "uploaded_at": ("-uploaded_at", "id"),
        "updated_at": ("-updated_at", "id"),
        "file_name": ("file_name", "id"),
    }
    return documents.order_by(
        *ordering_fields.get(ordering, ordering_fields["relevance"])
    )


def _fuzzy_document_results(documents, query: str, *, match_type: str, ordering: str):
    fuzzy_threshold = get_fuzzy_search_threshold(query)
    candidate_limit = int(
        getattr(settings, "DOCUMENT_SEARCH_FUZZY_CANDIDATE_LIMIT", 200)
    )
    matches = {}

    if match_type in {"all", "metadata"}:
        metadata_candidates = list(
            documents.annotate(
                filename_word_rank=TrigramWordSimilarity(query, "file_name"),
                filename_rank=TrigramSimilarity("file_name", query),
                document_number_word_rank=TrigramWordSimilarity(
                    query,
                    "full_document_number",
                ),
                identification_word_rank=TrigramWordSimilarity(
                    query,
                    "identification_number",
                ),
            )
            .annotate(
                candidate_rank=Greatest(
                    Coalesce("filename_word_rank", Value(0.0)),
                    Coalesce("filename_rank", Value(0.0)),
                    Coalesce("document_number_word_rank", Value(0.0)),
                    Coalesce("identification_word_rank", Value(0.0)),
                )
            )
            .filter(candidate_rank__gte=fuzzy_threshold)
            .order_by("-candidate_rank", "id")[:candidate_limit]
        )
        for document in metadata_candidates:
            score = calculate_fuzzy_token_score(document_metadata_text(document), query)
            if score > 0:
                matches[document.id] = {
                    "score": score,
                    "segment_id": None,
                    "metadata": True,
                }

    if match_type in {"all", "content"}:
        content_candidates = list(
            DocumentContentSegment.objects.filter(
                document__in=documents,
                index__status__in=[
                    DocumentSearchIndex.Status.SEARCHABLE,
                    DocumentSearchIndex.Status.PARTIAL,
                ],
            )
            .exclude(sequence=0)
            .annotate(
                content_word_rank=TrigramWordSimilarity(query, "text"),
            )
            .filter(content_word_rank__gte=fuzzy_threshold)
            .order_by("-content_word_rank", "id")[:candidate_limit]
        )
        for segment in content_candidates:
            score = calculate_fuzzy_token_score(segment.text, query)
            current = matches.get(segment.document_id)
            if score > 0 and (current is None or score > current["score"]):
                matches[segment.document_id] = {
                    "score": score,
                    "segment_id": segment.id,
                    "metadata": False,
                }

    if not matches:
        return documents.none().annotate(
            rank=Value(0.0, output_field=models.FloatField()),
            matched_content_segment_id=Value(
                None,
                output_field=models.IntegerField(),
            ),
            matched_by_metadata=Value(False, output_field=models.BooleanField()),
            match_quality=Value("fuzzy", output_field=models.CharField()),
        )

    rank_expression = Case(
        *[
            When(id=document_id, then=Value(match["score"]))
            for document_id, match in matches.items()
        ],
        default=Value(0.0),
        output_field=models.FloatField(),
    )
    segment_expression = Case(
        *[
            When(id=document_id, then=Value(match["segment_id"]))
            for document_id, match in matches.items()
        ],
        default=Value(None),
        output_field=models.IntegerField(),
    )
    metadata_expression = Case(
        *[
            When(id=document_id, then=Value(match["metadata"]))
            for document_id, match in matches.items()
        ],
        default=Value(False),
        output_field=models.BooleanField(),
    )
    results = documents.filter(id__in=matches).annotate(
        rank=rank_expression,
        matched_content_segment_id=segment_expression,
        matched_by_metadata=metadata_expression,
        match_quality=Value("fuzzy", output_field=models.CharField()),
    )
    return _ordered_document_results(results, ordering)


def search_documents(request, *, documents=None):
    query = (request.query_params.get("q") or "").strip()
    if len(query) < 3:
        return Document.objects.none(), 0

    documents = documents if documents is not None else accessible_document_queryset(request)
    folder_id = request.query_params.get("folder_id")
    if folder_id:
        documents = documents.filter(folder_id=folder_id)

    match_type = (request.query_params.get("match_type") or "all").lower()
    if match_type not in {"all", "content", "metadata"}:
        match_type = "all"

    ordering = (request.query_params.get("ordering") or "relevance").lower()
    search_query = SearchQuery(query, config="simple", search_type="websearch")
    content_segments = (
        DocumentContentSegment.objects.filter(
            document_id=OuterRef("pk"),
            index__status__in=[
                DocumentSearchIndex.Status.SEARCHABLE,
                DocumentSearchIndex.Status.PARTIAL,
            ],
        )
        .exclude(sequence=0)
        .annotate(
            search_rank=SearchRank(F("search_vector"), search_query),
            literal_rank=Case(
                When(
                    Q(text__icontains=query) | Q(text__unaccent__icontains=query),
                    then=Value(0.75),
                ),
                default=Value(0.0),
                output_field=models.FloatField(),
            ),
        )
        .filter(
            Q(search_vector=search_query)
            | Q(text__icontains=query)
            | Q(text__unaccent__icontains=query)
        )
        .annotate(
            result_rank=Greatest(
                Coalesce("search_rank", Value(0.0)),
                F("literal_rank"),
            )
        )
        .order_by("-result_rank", "sequence", "id")
    )

    content_rank = Value(0.0, output_field=models.FloatField())
    matched_segment_id = Value(None, output_field=models.IntegerField())
    if match_type in {"all", "content"}:
        content_rank = Coalesce(
            Subquery(
                content_segments.values("result_rank")[:1],
                output_field=models.FloatField(),
            ),
            Value(0.0),
        )
        matched_segment_id = Subquery(
            content_segments.values("id")[:1],
            output_field=models.IntegerField(),
        )

    metadata_rank = Value(0.0, output_field=models.FloatField())
    if match_type in {"all", "metadata"}:
        metadata_rank = build_document_metadata_rank(query)

    results = (
        documents.annotate(
            content_rank=content_rank,
            metadata_rank=metadata_rank,
            matched_content_segment_id=matched_segment_id,
        )
        .annotate(
            rank=Greatest(F("content_rank"), F("metadata_rank")),
            matched_by_metadata=Case(
                When(
                    Q(metadata_rank__gt=0)
                    & Q(metadata_rank__gte=F("content_rank")),
                    then=Value(True),
                ),
                default=Value(False),
                output_field=models.BooleanField(),
            ),
            match_quality=Value("exact", output_field=models.CharField()),
        )
        .filter(rank__gt=0)
    )

    if not results.exists():
        results = _fuzzy_document_results(
            documents,
            query,
            match_type=match_type,
            ordering=ordering,
        )
    else:
        results = _ordered_document_results(results, ordering)
    return results, results.count()


def get_document_indexing_summary(request, *, documents=None) -> dict:
    documents = documents if documents is not None else accessible_document_queryset(request)
    folder_id = request.query_params.get("folder_id")
    if folder_id:
        documents = documents.filter(folder_id=folder_id)

    counts = {
        row["status"]: row["count"]
        for row in DocumentSearchIndex.objects.filter(document__in=documents)
        .values("status")
        .annotate(count=models.Count("id"))
    }
    return {
        "pending": counts.get(DocumentSearchIndex.Status.PENDING, 0),
        "processing": counts.get(DocumentSearchIndex.Status.PROCESSING, 0),
        "failed": counts.get(DocumentSearchIndex.Status.FAILED, 0),
        "unsupported": counts.get(DocumentSearchIndex.Status.UNSUPPORTED, 0),
        "not_indexed": documents.filter(search_index__isnull=True).count(),
    }


def build_search_result(
    document: Document,
    query: str,
    *,
    segments_by_id=None,
    request=None,
) -> dict:
    folder = document.folder
    file_url = document.file.url if document.file else None
    if file_url and request is not None:
        file_url = request.build_absolute_uri(file_url)
    segment_id = getattr(document, "matched_content_segment_id", None)
    segment = (segments_by_id or {}).get(segment_id)
    matched_by_metadata = getattr(document, "matched_by_metadata", False)
    if matched_by_metadata or segment is None:
        matched_text = document_metadata_text(document)
        match_type = "metadata"
        location = {"type": "document", "number": 0, "label": "Metadata"}
    else:
        matched_text = segment.text
        match_type = "content"
        location = {
            "type": segment.segment_type,
            "number": segment.sequence,
            "label": segment.location_label or segment.get_segment_type_display(),
        }
    matched_terms = find_matching_terms(matched_text, query)
    try:
        index_status = document.search_index.status
    except DocumentSearchIndex.DoesNotExist:
        index_status = "not_indexed"
    return {
        "document_id": document.id,
        "file_name": document.file_name or os.path.basename(document.file.name),
        "folder": {"id": folder.id, "name": folder.name},
        "folder_breadcrumb": folder_breadcrumb(folder),
        "match_type": match_type,
        "snippet": make_snippet(
            matched_text,
            query,
            matched_terms=matched_terms,
        ),
        "matched_terms": matched_terms,
        "match_quality": getattr(document, "match_quality", "exact"),
        "location": location,
        "rank": float(document.rank or 0),
        "document_type": document.document_type,
        "status": document.status,
        "author": {
            "id": document.author_id,
            "name": document.author.full_name.strip() if document.author_id else None,
        },
        "uploaded_at": document.uploaded_at,
        "updated_at": document.updated_at,
        "index_status": index_status,
        "file_url": file_url,
    }


def folder_breadcrumb(folder) -> List[str]:
    names = []
    current = folder
    guard = 0
    while current and guard < 50:
        names.append(current.name)
        current = current.parent
        guard += 1
    return list(reversed(names))


SEARCH_WORD_PATTERN = re.compile(r"[^\W_]+(?:[-'][^\W_]+)*", flags=re.UNICODE)


def extract_search_terms(value: str) -> List[str]:
    terms = []
    for token in SEARCH_WORD_PATTERN.findall((value or "").casefold()):
        terms.extend(part for part in re.split(r"[-']", token) if part)
    return list(dict.fromkeys(terms))


def get_fuzzy_search_threshold(query: str) -> float:
    terms = extract_search_terms(query)
    if len(terms) > 1:
        return 0.45
    length = len(terms[0]) if terms else len(query.strip())
    if length <= 3:
        return 0.75
    if length == 4:
        return 0.68
    if length <= 7:
        return 0.50
    return 0.45


def get_token_match_threshold(token: str) -> float:
    if len(token) <= 3:
        return 0.86
    if len(token) <= 5:
        return 0.80
    if len(token) <= 7:
        return 0.78
    return 0.72


def calculate_fuzzy_token_score(text: str, query: str) -> float:
    safe_text = normalize_extracted_text(text)
    safe_query = normalize_extracted_text(query)
    query_terms = extract_search_terms(safe_query)
    if not safe_text or not query_terms:
        return 0.0

    candidate_terms = extract_search_terms(safe_text)
    if not candidate_terms:
        return 0.0

    term_scores = []
    for query_term in query_terms:
        best_score = max(
            (
                SequenceMatcher(None, query_term, candidate_term).ratio()
                for candidate_term in candidate_terms
                if abs(len(candidate_term) - len(query_term))
                <= max(2, len(query_term) // 3)
            ),
            default=0.0,
        )
        if best_score < get_token_match_threshold(query_term):
            return 0.0
        term_scores.append(best_score)

    return sum(term_scores) / len(term_scores)


def find_matching_terms(text: str, query: str, limit: int = 8) -> List[str]:
    safe_text = normalize_extracted_text(text)
    safe_query = normalize_extracted_text(query)
    if not safe_text or not safe_query:
        return []

    phrase_match = re.search(re.escape(safe_query), safe_text, flags=re.IGNORECASE)
    if phrase_match:
        return [phrase_match.group(0)]

    query_terms = extract_search_terms(safe_query)
    if not query_terms:
        return []

    candidates = []
    for match in SEARCH_WORD_PATTERN.finditer(safe_text):
        original_token = match.group(0)
        token_parts = re.split(r"[-']", original_token)
        candidates.extend(
            (part, part.casefold())
            for part in token_parts
            if part
        )
    matched_terms = []
    for query_term in query_terms[:limit]:
        best_term = ""
        best_score = 0.0
        for original_term, normalized_term in candidates:
            if normalized_term == query_term:
                best_term = original_term
                best_score = 1.0
                break
            if abs(len(normalized_term) - len(query_term)) > max(2, len(query_term) // 3):
                continue
            score = SequenceMatcher(None, query_term, normalized_term).ratio()
            if score > best_score:
                best_term = original_term
                best_score = score
        if best_term and best_score >= get_token_match_threshold(query_term):
            if best_term.casefold() not in {term.casefold() for term in matched_terms}:
                matched_terms.append(best_term)

    return matched_terms


def make_snippet(
    text: str,
    query: str,
    window: int = 90,
    *,
    matched_terms: Optional[List[str]] = None,
) -> str:
    safe_text = normalize_extracted_text(text)
    if not safe_text:
        return ""
    pattern = re.compile(re.escape(query), flags=re.IGNORECASE)
    match = pattern.search(safe_text)
    if not match:
        for term in matched_terms or find_matching_terms(safe_text, query):
            match = re.search(re.escape(term), safe_text, flags=re.IGNORECASE)
            if match:
                break
    if not match:
        return safe_text[: window * 2] + ("..." if len(safe_text) > window * 2 else "")
    start = max(0, match.start() - window)
    end = min(len(safe_text), match.end() + window)
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(safe_text) else ""
    return f"{prefix}{safe_text[start:end]}{suffix}"


def retry_due_jobs(limit: int = 100) -> int:
    job_ids = list(
        DocumentIndexJob.objects.filter(
            status__in=[DocumentIndexJob.Status.QUEUED, DocumentIndexJob.Status.RETRYING],
            available_at__lte=timezone.now(),
        ).values_list("id", flat=True)[:limit]
    )
    for job_id in job_ids:
        dispatch_document_indexing_job(job_id)
    return len(job_ids)
