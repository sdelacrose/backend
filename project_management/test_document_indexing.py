import tempfile
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZIP_STORED, ZipFile

from django.core.files.base import ContentFile
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from rest_framework.request import Request
from rest_framework.test import APIClient, APIRequestFactory

from models.models import User
from project_management.models import (
    Document,
    DocumentContentSegment,
    DocumentIndexAttempt,
    DocumentIndexJob,
    DocumentSearchIndex,
    Folder,
)
from project_management.services.document_indexing import (
    DocumentIndexingError,
    extract_cad_interchange,
    extract_docx,
    extract_dxf_text_values,
    extract_email,
    extract_email_attachment_text,
    extract_file_content,
    extract_iges_text_values,
    extract_pdf,
    extract_pptx,
    extract_step_text_values,
    extract_xlsx,
    ensure_extraction_path_suffix,
    find_matching_terms,
    make_snippet,
    normalize_extracted_text,
    normalize_extracted_lines,
    prepare_image_for_ocr,
    process_index_job,
    resolve_extraction_extension,
    read_pdf_segments,
    safe_tmp_filename,
    search_documents,
    split_search_vector_text,
)
from project_management.serializers.document import DocumentSerializer


class DocumentExtractorTests(SimpleTestCase):
    def test_normalized_text_removes_database_unsafe_null_bytes(self):
        self.assertEqual(
            normalize_extracted_text("Specification\x00reference\nA-100"),
            "Specification reference A-100",
        )

    def test_line_normalization_preserves_pdf_table_rows(self):
        self.assertEqual(
            normalize_extracted_lines(
                "Assembly\tPart number\r\nMotor Assembly\tM-100\n\nValve  Assembly\tV-200"
            ),
            "Assembly Part number\nMotor Assembly M-100\nValve Assembly V-200",
        )

    def test_pdf_layout_extraction_keeps_distinct_table_rows(self):
        from weasyprint import HTML

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "assembly-table.pdf"
            HTML(
                string=(
                    "<table><tr><th>Assembly</th><th>Part number</th></tr>"
                    "<tr><td>Motor Assembly</td><td>M-100</td></tr>"
                    "<tr><td>Valve Assembly</td><td>V-200</td></tr></table>"
                )
            ).write_pdf(pdf_path)

            segments, page_count = read_pdf_segments(pdf_path)

        self.assertEqual(page_count, 1)
        lines = segments[0].text.splitlines()
        self.assertTrue(any("Motor Assembly" in line and "M-100" in line for line in lines))
        self.assertTrue(any("Valve Assembly" in line and "V-200" in line for line in lines))

    def test_extraction_extension_falls_back_to_storage_name_and_detected_mime(self):
        self.assertEqual(
            resolve_extraction_extension(
                "Generated invoice",
                "invoices/record.pdf",
                "application/pdf",
            ),
            ".pdf",
        )
        self.assertEqual(
            resolve_extraction_extension(
                "Generated invoice",
                "stored-object",
                "application/pdf",
            ),
            ".pdf",
        )
        self.assertEqual(
            resolve_extraction_extension(
                "drawing.dwg",
                "stored-object",
                "application/octet-stream",
            ),
            ".dwg",
        )
        self.assertEqual(
            resolve_extraction_extension(
                "record.mp4",
                "stored-object",
                "video/mp4",
            ),
            "",
        )

    def test_extraction_path_preserves_detected_spreadsheet_suffix(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            original = Path(temp_dir) / "Test_OS"
            original.write_bytes(b"PK\x03\x04spreadsheet")

            extraction_path = ensure_extraction_path_suffix(original, ".xlsx")

            self.assertEqual(extraction_path.suffix, ".xlsx")
            self.assertEqual(extraction_path.read_bytes(), b"PK\x03\x04spreadsheet")
            self.assertFalse(original.exists())

    @patch("project_management.services.document_indexing.detect_mime_type")
    @patch("project_management.services.document_indexing.extract_xlsx")
    @patch("project_management.services.document_indexing.default_storage")
    def test_extensionless_display_name_uses_storage_file_extension(
        self,
        storage,
        extract_spreadsheet,
        detect_mime_type,
    ):
        storage.open.return_value = BytesIO(b"PK\x03\x04spreadsheet")
        detect_mime_type.return_value = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        extract_spreadsheet.return_value = ([], {"method": "native", "page_count": 0})
        file_field = SimpleNamespace(
            name="object_assessment/project_254/Documents/TopCustomers_6.xlsx"
        )

        extract_file_content(file_field, "Test OS")

        extraction_path = extract_spreadsheet.call_args.args[0]
        self.assertEqual(extraction_path.suffix, ".xlsx")

    @patch("project_management.services.document_indexing.detect_mime_type")
    @patch("project_management.services.document_indexing.default_storage")
    def test_zip_is_classified_as_stored_only_without_reading_archive_entries(
        self,
        storage,
        detect_mime_type,
    ):
        storage.open.return_value = BytesIO(
            b"PK\x03\x04archive payload that must not be extracted"
        )
        detect_mime_type.return_value = "application/zip"
        file_field = SimpleNamespace(name="documents/customer-demand.zip")

        segments, metadata = extract_file_content(
            file_field,
            "customer-demand.zip",
        )

        self.assertEqual(segments, [])
        self.assertTrue(metadata["stored_only"])
        self.assertEqual(
            metadata["extraction_method"],
            DocumentSearchIndex.ExtractionMethod.NONE,
        )
        self.assertEqual(metadata["mime_type"], "application/zip")

    def test_search_vector_text_is_split_without_losing_unicode_words(self):
        text = " ".join(["pressure", "Übertragung", "étanchéité"] * 20)

        chunks = split_search_vector_text(text, max_bytes=80)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.encode("utf-8")) <= 80 for chunk in chunks))
        self.assertEqual(" ".join(chunks).split(), text.split())

    def test_invalid_spreadsheet_raises_controlled_indexing_error(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "invalid.xlsx"
            path.write_bytes(b"not-an-xlsx-file")

            with self.assertRaises(DocumentIndexingError) as context:
                extract_xlsx(path, None)

        self.assertEqual(context.exception.category, "invalid_file")

    def test_fuzzy_snippet_uses_the_actual_document_word(self):
        text = "The operating history covers the twentieth century in detail."

        matched_terms = find_matching_terms(text, "centry")
        snippet = make_snippet(text, "centry", matched_terms=matched_terms)

        self.assertEqual(matched_terms, ["century"])
        self.assertIn("century", snippet)

    def test_office_and_pdf_extractors_preserve_locations(self):
        from docx import Document as DocxDocument
        from openpyxl import Workbook
        from pptx import Presentation
        from weasyprint import HTML

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)

            docx_path = temp_path / "specification.docx"
            docx = DocxDocument()
            docx.add_paragraph("Required pressure is eight bar")
            docx.sections[0].header.paragraphs[0].text = "Header reference HDR-100"
            docx.sections[0].footer.paragraphs[0].text = "Footer reference FTR-200"
            docx.save(docx_path)
            docx_segments, _ = extract_docx(docx_path, None)
            self.assertIn("eight bar", docx_segments[0].text)
            self.assertIn("HDR-100", docx_segments[0].text)
            self.assertIn("FTR-200", docx_segments[0].text)

            xlsx_path = temp_path / "parts.xlsx"
            workbook = Workbook()
            workbook.active.title = "Requirements"
            workbook.active.append(["Skill", "Drill"])
            workbook.create_sheet("Suppliers").append(["FESTO", "Cylinder"])
            workbook.save(xlsx_path)
            xlsx_segments, _ = extract_xlsx(xlsx_path, None)
            self.assertEqual(
                [segment.location_label for segment in xlsx_segments],
                ["Requirements", "Suppliers"],
            )

            pptx_path = temp_path / "handover.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Handover"
            slide.placeholders[1].text = "Hydraulic assembly"
            presentation.save(pptx_path)
            pptx_segments, _ = extract_pptx(pptx_path, None)
            self.assertEqual(pptx_segments[0].location_label, "Slide 1")
            self.assertIn("Hydraulic assembly", pptx_segments[0].text)

            pdf_path = temp_path / "proposal.pdf"
            HTML(string="<p>Unique proposal reference EC-4821</p>").write_pdf(pdf_path)
            pdf_segments, pdf_metadata = extract_pdf(pdf_path, "application/pdf")
            self.assertEqual(pdf_metadata["page_count"], 1)
            self.assertIn("EC-4821", pdf_segments[0].text)

    def test_email_body_and_supported_attachment_are_extracted(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "request.eml"
            message = EmailMessage()
            message["Subject"] = "Proposal request"
            message.set_content("Please quote the pneumatic assembly.")
            message.add_attachment(
                b"Attachment reference RFQ-9087",
                maintype="text",
                subtype="plain",
                filename="request.txt",
            )
            path.write_bytes(message.as_bytes())

            segments, _ = extract_email(path, ".eml", "message/rfc822")
            self.assertEqual(
                [segment.segment_type for segment in segments],
                ["email_body", "attachment"],
            )
            self.assertIn("RFQ-9087", segments[1].text)

    def test_email_cad_attachments_follow_the_same_search_rules(self):
        step_payload = (
            b"ISO-10303-21;#10=PRODUCT('P-100','Pump Housing',"
            b"'FESTO assembly',());END-ISO-10303-21;"
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            open_cad_text = extract_email_attachment_text(
                temp_path,
                "pump-housing.step",
                step_payload,
                "model/step",
            )
            proprietary_cad_text = extract_email_attachment_text(
                temp_path,
                "pump-housing.catpart",
                b"\x00\x01proprietary content",
                "application/octet-stream",
            )

        self.assertIn("Pump Housing", open_cad_text)
        self.assertEqual(proprietary_cad_text, "pump-housing.catpart")

    def test_safe_filename_and_snippet_are_plain_text(self):
        self.assertEqual(
            safe_tmp_filename("../../unsafe<script>.pdf"),
            "unsafe_script_.pdf",
        )
        snippet = make_snippet(
            "prefix <script>alert(1)</script> operating pressure eight bar suffix",
            "eight bar",
        )
        self.assertIn("<script>", snippet)
        self.assertIn("eight bar", snippet)

    def test_open_cad_interchange_parsers_extract_searchable_properties(self):
        dxf_text = "\n".join(
            [
                "0",
                "SECTION",
                "2",
                "ENTITIES",
                "0",
                "TEXT",
                "8",
                "SPECIFICATIONS",
                "1",
                "Operating pressure 8 bar",
            ]
        )
        step_text = (
            "ISO-10303-21;"
            "#10=PRODUCT('P-100','Pump Housing','FESTO assembly',());"
            "END-ISO-10303-21;"
        )
        iges_title = "Pump Assembly Search Fixture"
        iges_text = f"{iges_title:<72}S0000001\n13HPump Assembly"

        self.assertIn("Operating pressure 8 bar", extract_dxf_text_values(dxf_text))
        self.assertIn("Pump Housing", extract_step_text_values(step_text))
        self.assertIn(iges_title, extract_iges_text_values(iges_text))
        self.assertIn("Pump Assembly", extract_iges_text_values(iges_text))

    def test_open_cad_binary_content_falls_back_to_metadata_only(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "pump.step"
            path.write_bytes(b"\x00\x01\x02binary-cad-data")

            segments, metadata = extract_cad_interchange(
                path,
                ".step",
                "model/step",
            )

        self.assertEqual(segments, [])
        self.assertTrue(metadata["metadata_only"])
        self.assertTrue(metadata["partial"])

    def test_gif_ocr_uses_first_frame_png(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "drawing.gif"
            first_frame = Image.new("RGB", (30, 20), "white")
            second_frame = Image.new("RGB", (30, 20), "black")
            first_frame.save(
                path,
                save_all=True,
                append_images=[second_frame],
                format="GIF",
            )

            ocr_path = prepare_image_for_ocr(path)
            with Image.open(ocr_path) as image:
                self.assertEqual(ocr_path.suffix, ".png")
                self.assertEqual(image.mode, "RGB")


class DocumentIndexingPersistenceTests(TestCase):
    def setUp(self):
        self.notification_patcher = patch(
            "ecpmpapp.utils.should_notify_user_comprehensive",
            return_value=False,
        )
        self.notification_patcher.start()
        self.media_directory = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(
            DOCUMENT_INDEXING_ENABLED=True,
            DOCUMENT_CONTENT_SEARCH_ENABLED=True,
            DOCUMENT_OCR_ENABLED=False,
            MEDIA_ROOT=self.media_directory.name,
        )
        self.settings_override.enable()
        self.user = User.objects.create_user(
            email="document-indexer@example.com",
            password="password",
            first_name="Document",
            last_name="Indexer",
            user_type="freelancer",
        )
        self.folder = Folder.objects.create(
            freelancer=self.user,
            name="Specifications",
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def tearDown(self):
        self.notification_patcher.stop()
        self.settings_override.disable()
        self.media_directory.cleanup()

    def create_document(self, file_name, content):
        document = Document(
            folder=self.folder,
            author=self.user,
            created_by=self.user,
            file_name=file_name,
            document_type="Specification",
            status="Draft",
        )
        with self.captureOnCommitCallbacks(execute=True):
            document.file.save(file_name, ContentFile(content), save=True)
        return document

    def test_filename_update_preserves_storage_key_and_extracted_content(self):
        document = self.create_document(
            "original-specification.txt",
            b"Operating pressure is eight bar.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        stored_file_name = document.file.name
        original_job_count = DocumentIndexJob.objects.filter(
            document=document
        ).count()
        original_segment_ids = list(
            DocumentContentSegment.objects.filter(document=document).values_list(
                "id", flat=True
            )
        )
        serializer = DocumentSerializer(
            instance=document,
            data={"file_name": "Customer specification"},
            partial=True,
            context={"user": self.user},
        )
        self.assertTrue(serializer.is_valid(), serializer.errors)

        with patch(
            "project_management.services.search_indexing.enqueue_search_indexing"
        ) as enqueue_search_indexing:
            with self.captureOnCommitCallbacks(execute=True):
                serializer.save()

        document.refresh_from_db()
        self.assertEqual(document.file_name, "Customer specification.txt")
        self.assertEqual(document.file.name, stored_file_name)
        self.assertEqual(
            DocumentIndexJob.objects.filter(document=document).count(),
            original_job_count,
        )
        self.assertEqual(
            list(
                DocumentContentSegment.objects.filter(document=document).values_list(
                    "id", flat=True
                )
            ),
            original_segment_ids,
        )
        enqueue_search_indexing.assert_called_once_with("document", document.id)

    def test_upload_is_idempotent_and_creates_searchable_segments(self):
        document = self.create_document(
            "pump-specification.txt",
            b"Operating pressure is eight bar. Unique code PUMP-4455.",
        )
        self.assertEqual(
            DocumentIndexJob.objects.filter(document=document).count(),
            1,
        )

        job = DocumentIndexJob.objects.get(document=document)
        process_index_job(job.id)

        job.refresh_from_db()
        index = DocumentSearchIndex.objects.get(document=document)
        self.assertEqual(job.status, DocumentIndexJob.Status.COMPLETED)
        self.assertEqual(index.status, DocumentSearchIndex.Status.SEARCHABLE)
        self.assertEqual(
            list(
                DocumentContentSegment.objects.filter(document=document).values_list(
                    "sequence", flat=True
                )
            ),
            [0, 1],
        )

        from project_management.services.document_indexing import (
            enqueue_document_indexing,
        )

        self.assertIsNone(enqueue_document_indexing(document.id))
        self.assertEqual(
            DocumentIndexJob.objects.filter(document=document).count(),
            1,
        )

    def test_parser_upgrade_requeues_unchanged_document(self):
        document = self.create_document(
            "parser-upgrade.txt",
            b"Assembly A part A-100.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        index = DocumentSearchIndex.objects.get(document=document)
        index.parser_version = "document-search-v1"
        index.save(update_fields=["parser_version"])

        from project_management.services.document_indexing import (
            enqueue_document_indexing,
        )

        replacement = enqueue_document_indexing(document.id)

        self.assertIsNotNone(replacement)
        self.assertEqual(replacement.status, DocumentIndexJob.Status.QUEUED)

    def test_pdf_table_rows_remain_separate_in_stored_segments(self):
        from weasyprint import HTML

        pdf = HTML(
            string=(
                "<table><tr><th>Assembly</th><th>Part number</th></tr>"
                "<tr><td>Motor Assembly</td><td>M-100</td></tr>"
                "<tr><td>Valve Assembly</td><td>V-200</td></tr></table>"
            )
        ).write_pdf()
        document = self.create_document("assembly-table.pdf", pdf)

        process_index_job(DocumentIndexJob.objects.get(document=document).id)

        page_text = DocumentContentSegment.objects.get(
            document=document,
            segment_type=DocumentContentSegment.SegmentType.PAGE,
        ).text
        self.assertIn("Motor Assembly", page_text)
        self.assertIn("Valve Assembly", page_text)
        self.assertGreaterEqual(page_text.count("\n"), 2)

    def test_zip_upload_completes_as_stored_only_metadata(self):
        archive_sentinel = b"SECRET-ARCHIVE-CONTENT-MUST-NOT-BE-INDEXED"
        archive = BytesIO()
        with ZipFile(archive, "w", compression=ZIP_STORED) as zip_file:
            zip_file.writestr("private-specification.txt", archive_sentinel)
        document = self.create_document(
            "customer-demand.zip",
            archive.getvalue(),
        )

        job = DocumentIndexJob.objects.get(document=document)
        process_index_job(job.id)

        job.refresh_from_db()
        index = DocumentSearchIndex.objects.get(document=document)
        segment_text = " ".join(
            DocumentContentSegment.objects.filter(document=document).values_list(
                "text",
                flat=True,
            )
        )
        self.assertEqual(job.status, DocumentIndexJob.Status.COMPLETED)
        self.assertEqual(index.status, DocumentSearchIndex.Status.UNSUPPORTED)
        self.assertEqual(
            index.extraction_method,
            DocumentSearchIndex.ExtractionMethod.NONE,
        )
        self.assertEqual(index.mime_type, "application/zip")
        self.assertIn("customer-demand.zip", segment_text)
        self.assertNotIn(archive_sentinel.decode(), segment_text)

    def test_xml_and_open_cad_content_are_searchable(self):
        xml_document = self.create_document(
            "specification.xml",
            b"<specification><manufacturer>FESTO</manufacturer>"
            b"<pressure>8 bar</pressure></specification>",
        )
        step_document = self.create_document(
            "pump-housing.step",
            b"ISO-10303-21;#10=PRODUCT('P-100','Pump Housing',"
            b"'FESTO assembly',());END-ISO-10303-21;",
        )
        for job in DocumentIndexJob.objects.filter(
            document__in=[xml_document, step_document]
        ):
            process_index_job(job.id)

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "FESTO assembly"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(
                id__in=[xml_document.id, step_document.id]
            ),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        self.assertEqual(results.first().id, step_document.id)

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "pressure 8 bar"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=xml_document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        self.assertEqual(results.first().id, xml_document.id)

    def test_proprietary_cad_is_metadata_searchable_only(self):
        document = self.create_document(
            "CATIA-pump-housing.catpart",
            b"\x00\x01hidden proprietary drawing content",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)

        index = DocumentSearchIndex.objects.get(document=document)
        segments = DocumentContentSegment.objects.filter(document=document)
        self.assertEqual(index.status, DocumentSearchIndex.Status.PARTIAL)
        self.assertEqual(list(segments.values_list("sequence", flat=True)), [0])

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "CATIA pump"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        result = results.first()
        self.assertEqual(result.id, document.id)
        self.assertTrue(result.matched_by_metadata)
        self.assertIsNone(result.matched_content_segment_id)

    def test_search_returns_one_best_result_and_respects_accessible_queryset(self):
        visible = self.create_document(
            "visible-specification.txt",
            b"The required coupling torque is 42 newton metres.",
        )
        hidden = self.create_document(
            "hidden-specification.txt",
            b"The required coupling torque is 42 newton metres.",
        )
        for job in DocumentIndexJob.objects.filter(document__in=[visible, hidden]):
            process_index_job(job.id)

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "coupling torque"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=visible.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        self.assertEqual(results.first().id, visible.id)

    def test_content_search_api_is_paginated_and_does_not_leak_other_freelancer_documents(self):
        first = self.create_document(
            "first-coupling-specification.txt",
            b"The required coupling torque is 42 newton metres.",
        )
        second = self.create_document(
            "second-coupling-specification.txt",
            b"The required coupling torque is 38 newton metres.",
        )
        other_user = User.objects.create_user(
            email="other-document-user@example.com",
            password="password",
            first_name="Other",
            last_name="User",
            user_type="freelancer",
        )
        other_folder = Folder.objects.create(
            freelancer=other_user,
            name="Private specifications",
        )
        hidden = Document(
            folder=other_folder,
            author=other_user,
            created_by=other_user,
            file_name="hidden-coupling-specification.txt",
            document_type="Specification",
            status="Draft",
        )
        with self.captureOnCommitCallbacks(execute=True):
            hidden.file.save(
                "hidden-coupling-specification.txt",
                ContentFile(b"The required coupling torque is 50 newton metres."),
                save=True,
            )

        for job in DocumentIndexJob.objects.filter(document__in=[first, second, hidden]):
            process_index_job(job.id)

        with patch(
            "project_management.services.document_access.require_any_permission_multiple_modules",
            return_value=False,
        ):
            response = self.client.get(
                "/api/project_management/document/content-search",
                {
                    "q": "coupling torque",
                    "status": "Draft",
                    "page": 1,
                    "page_size": 1,
                },
            )
            hidden_status = self.client.get(
                f"/api/project_management/document/{hidden.id}/index-status"
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 2)
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(response.data["next"], 2)
        self.assertIn(
            response.data["results"][0]["document_id"],
            {first.id, second.id},
        )
        self.assertEqual(response.data["results"][0]["match_quality"], "exact")
        self.assertIn("coupling torque", response.data["results"][0]["matched_terms"])
        self.assertEqual(hidden_status.status_code, 404)

    def test_filename_typo_can_match_without_exposing_unrelated_content(self):
        document = self.create_document(
            "pneumatic-specification.txt",
            b"The internal test phrase is not part of the query.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "pneumatic specifcation"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        self.assertEqual(results.first().id, document.id)

    def test_content_typo_matches_the_extracted_document_word(self):
        document = self.create_document(
            "historical-specification.txt",
            b"The specification covers machinery from the twentieth century.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "centry"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        result = results.first()
        self.assertEqual(count, 1)
        self.assertEqual(result.id, document.id)
        self.assertEqual(result.match_quality, "fuzzy")
        segment = DocumentContentSegment.objects.get(
            id=result.matched_content_segment_id,
        )
        self.assertEqual(find_matching_terms(segment.text, "centry"), ["century"])

    def test_unrelated_content_is_not_returned_for_a_typo_query(self):
        document = self.create_document(
            "pressure-specification.txt",
            b"The pneumatic assembly requires an operating pressure of eight bar.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "centry"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 0)
        self.assertFalse(results.exists())

    def test_typo_reranking_prefers_the_intended_word_over_a_weaker_match(self):
        target = self.create_document(
            "historical-period.txt",
            b"The manuscript dates from the seventh century.",
        )
        decoy = self.create_document(
            "central-office.txt",
            b"The central office controls the regional workflow.",
        )
        for job in DocumentIndexJob.objects.filter(document__in=[target, decoy]):
            process_index_job(job.id)

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "centry"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id__in=[target.id, decoy.id]),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 1)
        self.assertEqual(results.first().id, target.id)

    def test_unindexed_document_is_searchable_by_filename(self):
        document = self.create_document(
            "Legacy_Invoice_Project_104.pdf",
            b"This document has not been indexed yet.",
        )
        DocumentSearchIndex.objects.filter(document=document).delete()
        self.assertFalse(
            DocumentSearchIndex.objects.filter(document=document).exists()
        )

        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "legacy invoice project 104"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        result = results.first()
        self.assertEqual(count, 1)
        self.assertEqual(result.id, document.id)
        self.assertTrue(result.matched_by_metadata)
        self.assertEqual(result.match_quality, "exact")

    def test_unindexed_filename_supports_typo_matching(self):
        document = self.create_document(
            "Pneumatic_Valve_Catalogue.pdf",
            b"This document has not been indexed yet.",
        )
        DocumentSearchIndex.objects.filter(document=document).delete()
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "pnumatic valve catalgue"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        result = results.first()
        self.assertEqual(count, 1)
        self.assertEqual(result.id, document.id)
        self.assertTrue(result.matched_by_metadata)
        self.assertEqual(result.match_quality, "fuzzy")

    def test_content_only_search_does_not_treat_filename_as_content(self):
        document = self.create_document(
            "Filename_Only_Reference.pdf",
            b"This document has not been indexed yet.",
        )
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "filename only reference", "match_type": "content"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 0)
        self.assertFalse(results.exists())

    def test_punctuation_only_query_does_not_match_every_document(self):
        document = self.create_document(
            "ordinary-reference.pdf",
            b"This document has not been indexed yet.",
        )
        request = Request(
            APIRequestFactory().get(
                "/api/project_management/document/content-search",
                {"q": "---"},
            )
        )
        with patch(
            "project_management.services.document_indexing.accessible_document_queryset",
            return_value=Document.objects.filter(id=document.id),
        ):
            results, count = search_documents(request)

        self.assertEqual(count, 0)
        self.assertFalse(results.exists())

    def test_file_replacement_removes_stale_content_before_reindexing(self):
        document = self.create_document(
            "replaceable-specification.txt",
            b"Original unique content OLD-REFERENCE-100.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        self.assertTrue(
            DocumentContentSegment.objects.filter(
                document=document,
                text__contains="OLD-REFERENCE-100",
            ).exists()
        )

        with self.captureOnCommitCallbacks(execute=True):
            document.file.save(
                "replaceable-specification.txt",
                ContentFile(b"Replacement unique content NEW-REFERENCE-200."),
                save=True,
            )

        self.assertFalse(
            DocumentContentSegment.objects.filter(document=document).exists()
        )
        replacement_job = DocumentIndexJob.objects.filter(
            document=document,
            status=DocumentIndexJob.Status.QUEUED,
        ).latest("id")
        process_index_job(replacement_job.id)
        self.assertFalse(
            DocumentContentSegment.objects.filter(
                document=document,
                text__contains="OLD-REFERENCE-100",
            ).exists()
        )
        self.assertTrue(
            DocumentContentSegment.objects.filter(
                document=document,
                text__contains="NEW-REFERENCE-200",
            ).exists()
        )

    def test_transient_failure_is_recorded_and_scheduled_for_retry(self):
        document = self.create_document(
            "temporary-failure.txt",
            b"Content that will encounter a temporary extraction failure.",
        )
        job = DocumentIndexJob.objects.get(document=document)
        with (
            patch(
                "project_management.services.document_indexing.extract_document",
                side_effect=DocumentIndexingError("Temporary parser failure.", "timeout"),
            ),
            patch(
                "project_management.services.document_indexing.dispatch_document_indexing_job"
            ) as dispatch,
        ):
            process_index_job(job.id)

        job.refresh_from_db()
        index = DocumentSearchIndex.objects.get(document=document)
        self.assertEqual(job.status, DocumentIndexJob.Status.RETRYING)
        self.assertEqual(index.status, DocumentSearchIndex.Status.PENDING)
        self.assertEqual(index.failure_category, "timeout")
        self.assertEqual(
            DocumentIndexAttempt.objects.get(job=job).failure_category,
            "timeout",
        )
        dispatch.assert_called_once_with(job.id)

    @override_settings(DOCUMENT_INDEXING_MAX_FILE_SIZE_MB=0)
    def test_oversized_file_fails_without_retrying(self):
        document = self.create_document(
            "oversized.txt",
            b"Any non-empty content exceeds a zero-byte test limit.",
        )
        job = DocumentIndexJob.objects.get(document=document)
        with patch(
            "project_management.services.document_indexing.dispatch_document_indexing_job"
        ) as dispatch:
            process_index_job(job.id)

        job.refresh_from_db()
        index = DocumentSearchIndex.objects.get(document=document)
        self.assertEqual(job.status, DocumentIndexJob.Status.FAILED)
        self.assertEqual(job.failure_category, "oversized")
        self.assertEqual(index.status, DocumentSearchIndex.Status.FAILED)
        dispatch.assert_not_called()

    def test_index_status_hides_parser_error_from_non_staff_user(self):
        document = self.create_document(
            "failed-index.txt",
            b"Content that is marked as failed for API visibility testing.",
        )
        index = DocumentSearchIndex.objects.get(document=document)
        index.status = DocumentSearchIndex.Status.FAILED
        index.error = "Internal parser details"
        index.failure_category = "extraction_error"
        index.save()

        with patch(
            "project_management.services.document_access.require_any_permission_multiple_modules",
            return_value=False,
        ):
            response = self.client.get(
                f"/api/project_management/document/{document.id}/index-status"
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["error"], "")
        self.assertEqual(response.data["failure_category"], "extraction_error")

    def test_retry_endpoint_requeues_a_failed_document(self):
        document = self.create_document(
            "failed-retry.txt",
            b"Content that can be retried after a transient parser failure.",
        )
        index = DocumentSearchIndex.objects.get(document=document)
        index.status = DocumentSearchIndex.Status.FAILED
        index.failure_category = "timeout"
        index.save(update_fields=["status", "failure_category", "updated_at"])

        with (
            patch(
                "project_management.services.document_access.require_any_permission_multiple_modules",
                return_value=False,
            ),
            patch(
                "project_management.views.document_search.DocumentPermission.has_update_permission",
                return_value=True,
            ),
            patch(
                "project_management.services.document_indexing.dispatch_document_indexing_job"
            ),
        ):
            response = self.client.post(
                f"/api/project_management/document/{document.id}/retry-indexing"
            )

        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.data["status"], DocumentIndexJob.Status.QUEUED)
        index.refresh_from_db()
        self.assertEqual(index.status, DocumentSearchIndex.Status.PENDING)

    def test_deleting_document_removes_search_state(self):
        document = self.create_document(
            "delete-index.txt",
            b"Search state must be removed with the source document.",
        )
        process_index_job(DocumentIndexJob.objects.get(document=document).id)
        document_id = document.id

        document.delete()

        self.assertFalse(DocumentSearchIndex.objects.filter(document_id=document_id).exists())
        self.assertFalse(DocumentContentSegment.objects.filter(document_id=document_id).exists())
        self.assertFalse(DocumentIndexJob.objects.filter(document_id=document_id).exists())

    def test_backfill_redispatches_due_jobs_before_queuing_new_jobs(self):
        document = self.create_document(
            "backfill-order.txt",
            b"The backfill command must not redispatch a newly queued job.",
        )
        events = []

        with (
            patch(
                "project_management.management.commands.backfill_document_index.retry_due_jobs",
                side_effect=lambda **_: events.append("retry") or 0,
            ),
            patch(
                "project_management.management.commands.backfill_document_index.enqueue_document_indexing",
                side_effect=lambda *_, **__: events.append("enqueue") or True,
            ),
        ):
            call_command(
                "backfill_document_index",
                limit=1,
                after_id=document.id - 1,
            )

        self.assertEqual(events, ["retry", "enqueue"])
