import hashlib
import io
import json
import os
import shutil
import subprocess
import tempfile
from datetime import date

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max
from docx import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from project_management.models import Document

from .audit import record_project_content_event
from .folders import get_project_content_folder
from .models import (
    GeneratedProjectDocumentRevision,
    ProjectDocumentGenerationRun,
)
from .template_validation import CONTENT_MARKER


class DocumentRenderingError(RuntimeError):
    pass


DOCUMENT_RENDERER_VERSION = "4"


def document_template_version(template):
    base_version = (
        f"company-{template.id}-v{template.version}" if template else "system-v1"
    )
    return f"{base_version}-renderer-v{DOCUMENT_RENDERER_VERSION}"


def _paragraph_before(paragraph, style=None):
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    created = Paragraph(element, paragraph._parent)
    if style:
        try:
            created.style = style
        except KeyError:
            pass
    return created


def _insert_paragraph_before(paragraph, text="", style=None):
    created = _paragraph_before(paragraph, style)
    if text:
        created.add_run(str(text))
    return created


def _render_generic_value(marker, value, level=0):
    if isinstance(value, str):
        _insert_paragraph_before(marker, value)
        return
    if isinstance(value, (int, float, bool)):
        _insert_paragraph_before(marker, str(value))
        return
    if isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                title = (
                    item.get("title")
                    or item.get("name")
                    or item.get("requirement")
                    or item.get("text")
                    or item.get("coverage")
                )
                if title:
                    _insert_paragraph_before(marker, title, "List Bullet")
                remainder = {
                    key: nested
                    for key, nested in item.items()
                    if key
                    not in {
                        "title",
                        "name",
                        "requirement",
                        "text",
                        "coverage",
                        "evidence",
                        "source_references",
                        "unresolved",
                    }
                }
                if remainder:
                    _render_generic_value(marker, remainder, level + 1)
            else:
                _insert_paragraph_before(marker, str(item), "List Bullet")
        return
    if isinstance(value, dict):
        for key, nested in value.items():
            if nested in (None, "", [], {}):
                continue
            _insert_paragraph_before(
                marker,
                key.replace("_", " ").title(),
                "Heading 3" if level else "Heading 2",
            )
            _render_generic_value(marker, nested, level + 1)


SECTION_ORDERS = {
    "offer": [
        "executive_summary",
        "customer_need",
        "proposed_scope",
        "proposed_solution",
        "deliverables",
        "milestones_and_schedule",
        "customer_responsibilities",
        "supplier_responsibilities",
        "commercial_values",
        "attachment_manifest",
        "assumptions",
        "exclusions",
        "requirements_coverage",
        "open_questions",
    ],
    "requirements_specification": [
        "purpose_and_scope",
        "stakeholders",
        "system_overview",
        "product_requirements",
        "functional_requirements",
        "performance_requirements",
        "non_functional_requirements",
        "interfaces",
        "constraints",
        "safety_and_compliance",
        "acceptance_conditions",
        "assumptions",
        "selected_workstations",
        "workstation_hierarchy",
        "evidence_metadata",
        "unresolved_decisions",
        "traceability_matrix",
        "unresolved_questions",
        "omitted_requirement_ids",
    ],
}


LABELS = {
    "en": {
        "product_requirements": "Product Requirements",
        "document": "Document",
        "document_type": "Type",
        "technical_specification": "Technical Specification",
        "executive_summary": "Executive Summary",
        "customer_need": "Customer Need",
        "proposed_scope": "Proposed Scope",
        "proposed_solution": "Proposed Solution",
        "deliverables": "Deliverables",
        "milestones_and_schedule": "Milestones and Schedule",
        "customer_responsibilities": "Customer Responsibilities",
        "supplier_responsibilities": "Supplier Responsibilities",
        "commercial_values": "Commercial Information",
        "attachment_manifest": "Attachments",
        "assumptions": "Assumptions",
        "exclusions": "Exclusions",
        "requirements_coverage": "Requirements Coverage",
        "open_questions": "Open Questions",
        "purpose_and_scope": "Purpose and Scope",
        "selected_workstations": "Selected Workstations",
        "workstation_hierarchy": "Workstation Hierarchy",
        "evidence_metadata": "Evidence and Version Metadata",
        "unresolved_decisions": "Unresolved Decisions",
        "stakeholders": "Stakeholders",
        "system_overview": "Product or System Overview",
        "functional_requirements": "Functional Requirements",
        "performance_requirements": "Performance Requirements",
        "non_functional_requirements": "Non-functional Requirements",
        "interfaces": "Interfaces",
        "constraints": "Constraints",
        "safety_and_compliance": "Safety and Compliance",
        "acceptance_conditions": "Acceptance Conditions",
        "traceability_matrix": "Requirements Traceability",
        "unresolved_questions": "Unresolved Questions",
        "omitted_requirement_ids": "Requirements Requiring Review",
        "document_information": "Document Information",
        "project": "Project",
        "revision": "Revision",
        "date": "Date",
        "status": "Status",
        "draft": "Draft",
        "approved": "Approved",
        "superseded": "Superseded",
        "unresolved": "Unresolved",
        "source_references": "Source References",
        "reference": "Reference",
        "source": "Source",
        "location": "Location",
        "excerpt": "Evidence Excerpt",
        "requirement": "Requirement",
        "coverage": "Coverage",
        "sources": "Sources",
    },
    "de": {
        "product_requirements": "Produktanforderungen",
        "document": "Dokument",
        "document_type": "Typ",
        "technical_specification": "Technische Spezifikation",
        "executive_summary": "Zusammenfassung",
        "customer_need": "Kundenbedarf",
        "proposed_scope": "Vorgeschlagener Leistungsumfang",
        "proposed_solution": "Vorgeschlagene Lösung",
        "deliverables": "Liefergegenstände",
        "milestones_and_schedule": "Meilensteine und Zeitplan",
        "customer_responsibilities": "Verantwortlichkeiten des Kunden",
        "supplier_responsibilities": "Verantwortlichkeiten des Lieferanten",
        "commercial_values": "Kaufmännische Angaben",
        "attachment_manifest": "Anhänge",
        "assumptions": "Annahmen",
        "exclusions": "Abgrenzungen",
        "requirements_coverage": "Anforderungsabdeckung",
        "open_questions": "Offene Fragen",
        "purpose_and_scope": "Zweck und Umfang",
        "selected_workstations": "Ausgewählte Arbeitsstationen",
        "workstation_hierarchy": "Arbeitsstationshierarchie",
        "evidence_metadata": "Nachweis- und Versionsmetadaten",
        "unresolved_decisions": "Offene Entscheidungen",
        "stakeholders": "Stakeholder",
        "system_overview": "Produkt- oder Systemübersicht",
        "functional_requirements": "Funktionale Anforderungen",
        "performance_requirements": "Leistungsanforderungen",
        "non_functional_requirements": "Nichtfunktionale Anforderungen",
        "interfaces": "Schnittstellen",
        "constraints": "Randbedingungen",
        "safety_and_compliance": "Sicherheit und Konformität",
        "acceptance_conditions": "Abnahmebedingungen",
        "traceability_matrix": "Anforderungsrückverfolgbarkeit",
        "unresolved_questions": "Ungeklärte Fragen",
        "omitted_requirement_ids": "Zu prüfende Anforderungen",
        "document_information": "Dokumentinformationen",
        "project": "Projekt",
        "revision": "Revision",
        "date": "Datum",
        "status": "Status",
        "draft": "Entwurf",
        "approved": "Freigegeben",
        "superseded": "Ersetzt",
        "unresolved": "Ungeklärt",
        "source_references": "Quellenverweise",
        "reference": "Referenz",
        "source": "Quelle",
        "location": "Fundstelle",
        "excerpt": "Belegauszug",
        "requirement": "Anforderung",
        "coverage": "Abdeckung",
        "sources": "Quellen",
    },
    "fr": {
        "product_requirements": "Exigences produit",
        "document": "Document",
        "document_type": "Type",
        "technical_specification": "Spécification technique",
        "executive_summary": "Résumé exécutif",
        "customer_need": "Besoin du client",
        "proposed_scope": "Périmètre proposé",
        "proposed_solution": "Solution proposée",
        "deliverables": "Livrables",
        "milestones_and_schedule": "Jalons et calendrier",
        "customer_responsibilities": "Responsabilités du client",
        "supplier_responsibilities": "Responsabilités du fournisseur",
        "commercial_values": "Informations commerciales",
        "attachment_manifest": "Pièces jointes",
        "assumptions": "Hypothèses",
        "exclusions": "Exclusions",
        "requirements_coverage": "Couverture des exigences",
        "open_questions": "Questions ouvertes",
        "purpose_and_scope": "Objectif et périmètre",
        "selected_workstations": "Postes de travail sélectionnés",
        "workstation_hierarchy": "Hiérarchie des postes de travail",
        "evidence_metadata": "Métadonnées de preuve et de version",
        "unresolved_decisions": "Décisions en suspens",
        "stakeholders": "Parties prenantes",
        "system_overview": "Vue d’ensemble du produit ou système",
        "functional_requirements": "Exigences fonctionnelles",
        "performance_requirements": "Exigences de performance",
        "non_functional_requirements": "Exigences non fonctionnelles",
        "interfaces": "Interfaces",
        "constraints": "Contraintes",
        "safety_and_compliance": "Sécurité et conformité",
        "acceptance_conditions": "Conditions d’acceptation",
        "traceability_matrix": "Traçabilité des exigences",
        "unresolved_questions": "Questions non résolues",
        "omitted_requirement_ids": "Exigences à examiner",
        "document_information": "Informations du document",
        "project": "Projet",
        "revision": "Révision",
        "date": "Date",
        "status": "Statut",
        "draft": "Brouillon",
        "approved": "Approuvé",
        "superseded": "Remplacé",
        "unresolved": "Non résolu",
        "source_references": "Références des sources",
        "reference": "Référence",
        "source": "Source",
        "location": "Emplacement",
        "excerpt": "Extrait justificatif",
        "requirement": "Exigence",
        "coverage": "Couverture",
        "sources": "Sources",
    },
}


def _label(locale, key):
    locale = locale if locale in LABELS else "en"
    return LABELS[locale].get(key, LABELS["en"].get(key, key.replace("_", " ").title()))


def _insert_table_before(document, marker, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Light Shading Accent 1"
    except KeyError:
        table.style = "Table Grid"
    table.autofit = not bool(widths)
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = width
    header_row = table.rows[0]
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
        if widths:
            cell.width = widths[index]
    tr_pr = header_row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    tr_pr.append(OxmlElement("w:cantSplit"))
    for values in rows:
        row = table.add_row()
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        cells = row.cells
        for index, value in enumerate(values):
            cells[index].text = str(value or "")
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[index].width = widths[index]
    marker._p.addprevious(table._tbl)
    return table


def _insert_key_value_table_before(document, marker, rows, widths=None):
    table = document.add_table(rows=0, cols=2)
    try:
        table.style = "Light Shading Accent 1"
    except KeyError:
        table.style = "Table Grid"
    table.autofit = not bool(widths)
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = width
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value or "")
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
        for index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = widths[index]
    marker._p.addprevious(table._tbl)
    return table


def _source_maps(snapshot):
    snapshot = snapshot or {}
    return {
        "document": {
            str(row.get("id")): row.get("name") or f"Document {row.get('id')}"
            for row in snapshot.get("documents", [])
        },
        "requirement": {
            str(row.get("id")): " - ".join(
                value for value in [row.get("object_number"), row.get("name")] if value
            )
            for row in snapshot.get("requirements", [])
        },
        "process_station": {
            str(row.get("id")): " - ".join(
                value for value in [row.get("object_number"), row.get("name")] if value
            )
            for row in snapshot.get("process_stations", [])
        },
        "customer_demand": {
            str(row.get("id")): " - ".join(
                value for value in [row.get("demand_number"), row.get("name")] if value
            )
            for row in snapshot.get("customer_demands", [])
        },
        "solution_space": {
            str(row.get("id")): row.get("name") or str(row.get("id"))
            for row in snapshot.get("solution_space", [])
        },
        "bill_of_material": {
            str(row.get("id")): " - ".join(
                value for value in [row.get("part_number"), row.get("name")] if value
            )
            for row in snapshot.get("bill_of_material", [])
        },
        "opportunity": {
            str(row.get("id")): row.get("title") or f"Opportunity {row.get('id')}"
            for row in snapshot.get("opportunities", [])
        },
        "project_workstation": {
            str(row.get("id")): row.get("name") or f"Workstation {row.get('id')}"
            for row in snapshot.get("project_workstations", [])
        },
    }


def _evidence_source_name(evidence, snapshot, maps):
    source_type = str(evidence.get("source_type") or "document")
    if source_type == "document":
        return maps["document"].get(
            str(evidence.get("document_id")),
            "Project document",
        )
    if source_type == "project":
        return (snapshot or {}).get("project", {}).get("name") or "Project record"
    source_id = evidence.get("source_id") or evidence.get("requirement_id")
    return maps.get(source_type, {}).get(
        str(source_id),
        source_type.replace("_", " ").title(),
    )


def _reference_key(evidence):
    return (
        str(evidence.get("source_type") or "document"),
        str(evidence.get("document_id") or ""),
        str(evidence.get("segment_id") or ""),
        str(evidence.get("source_id") or ""),
        str(evidence.get("requirement_id") or ""),
        str(evidence.get("location") or ""),
        str(evidence.get("quote") or ""),
    )


def _register_references(evidence_rows, references, snapshot, maps):
    labels = []
    for evidence in evidence_rows or []:
        if not isinstance(evidence, dict):
            continue
        quote = str(evidence.get("quote") or "").strip()
        if not quote or quote.casefold() in {"null", "none", "n/a"}:
            continue
        key = _reference_key(evidence)
        if key not in references:
            if len(quote) > 320:
                quote = f"{quote[:317].rstrip()}..."
            source_type = str(evidence.get("source_type") or "document")
            location = str(evidence.get("location") or "").strip()
            if source_type == "document" and location.casefold() == "document":
                location = ""
            elif source_type in {"project", "opportunity"}:
                location = location.replace("_", " ").title()
            elif source_type != "document":
                location = ""
            references[key] = {
                "label": f"S{len(references) + 1}",
                "source": _evidence_source_name(evidence, snapshot, maps),
                "location": location,
                "quote": quote,
            }
        labels.append(references[key]["label"])
    return labels


def _insert_statement(
    marker, statement, references, snapshot, maps, locale, hierarchy_level=None
):
    if isinstance(statement, str):
        text = statement.strip()
        unresolved = False
        evidence = []
    elif isinstance(statement, dict):
        text = str(statement.get("text") or statement.get("coverage") or "").strip()
        unresolved = bool(statement.get("unresolved"))
        evidence = statement.get("evidence") or []
    else:
        text = str(statement).strip()
        unresolved = False
        evidence = []
    if not text:
        return
    paragraph = _paragraph_before(marker, "List Bullet")
    if hierarchy_level is not None:
        paragraph.paragraph_format.left_indent = Inches(
            0.25 + 0.25 * max(1, int(hierarchy_level))
        )
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    if unresolved:
        prefix = paragraph.add_run(f"{_label(locale, 'unresolved')}: ")
        prefix.bold = True
    paragraph.add_run(text)
    reference_labels = _register_references(evidence, references, snapshot, maps)
    if reference_labels:
        source_run = paragraph.add_run(f" [{', '.join(reference_labels)}]")
        source_run.italic = True
        source_run.font.size = Pt(8)


def _requirement_display(requirement_id, maps):
    return maps.get("requirement", {}).get(str(requirement_id), str(requirement_id))


def _insert_document_information(document, marker, snapshot, locale, revision):
    project = (snapshot or {}).get("project") or {}
    if not project and revision is None:
        return
    _insert_paragraph_before(marker, _label(locale, "document_information"), "Heading 1")
    rows = []
    if project.get("name"):
        rows.append((_label(locale, "project"), project["name"]))
    if revision is not None:
        rows.append((_label(locale, "revision"), f"R{revision}"))
    rows.extend(
        [
            (_label(locale, "date"), date.today().isoformat()),
            (_label(locale, "status"), _label(locale, "draft")),
        ]
    )
    _insert_key_value_table_before(
        document,
        marker,
        rows,
        [Inches(1.5), Inches(5.0)],
    )


def _render_structured_content(
    document,
    marker,
    kind,
    content,
    locale,
    source_snapshot=None,
):
    references = {}
    maps = _source_maps(source_snapshot)
    rendered_sections = 0
    for key in SECTION_ORDERS[kind]:
        value = content.get(key)
        if value in (None, "", [], {}):
            continue
        _insert_paragraph_before(marker, _label(locale, key), "Heading 1")
        rendered_sections += 1
        if key == "acceptance_conditions":
            # Keep execution prerequisites beside the criteria. Product targets
            # remain valid even when their measurement method is not yet agreed.
            # Repeat the source's open decisions, not an inferred engineering rule.
            for decision in content.get("unresolved_decisions") or content.get("unresolved_questions") or []:
                statement = {"text": decision} if isinstance(decision, str) else dict(decision)
                statement["unresolved"] = True
                _insert_statement(
                    marker, statement, references, source_snapshot, maps, locale
                )
        if key in {"requirements_coverage", "traceability_matrix"}:
            rows = []
            for row in value:
                labels = _register_references(
                    row.get("evidence") or [], references, source_snapshot, maps
                )
                rows.append(
                    (
                        _requirement_display(row.get("requirement_id"), maps),
                        row.get("coverage") or "",
                        ", ".join(labels),
                    )
                )
            _insert_table_before(
                document,
                marker,
                [
                    _label(locale, "requirement"),
                    _label(locale, "coverage"),
                    _label(locale, "sources"),
                ],
                rows,
                [Inches(1.8), Inches(3.7), Inches(1.1)],
            )
        elif key == "attachment_manifest":
            rows = [
                (
                    row.get("name") or "",
                    _label(locale, row["document_type"]) if row.get("document_type") else "",
                    row.get("version") or "",
                    _label(locale, row["generation_status"]) if row.get("generation_status") else "",
                )
                for row in value
            ]
            _insert_table_before(
                document,
                marker,
                [
                    _label(locale, "document"),
                    _label(locale, "document_type"),
                    _label(locale, "revision"),
                    _label(locale, "status"),
                ],
                rows,
                [Inches(2.95), Inches(1.35), Inches(1.0), Inches(1.2)],
            )
        elif key == "omitted_requirement_ids":
            for requirement_id in value:
                _insert_statement(
                    marker,
                    {
                        "text": _requirement_display(requirement_id, maps),
                        "unresolved": True,
                    },
                    references,
                    source_snapshot,
                    maps,
                    locale,
                )
        else:
            for statement in value if isinstance(value, list) else [value]:
                _insert_statement(
                    marker,
                    statement,
                    references,
                    source_snapshot,
                    maps,
                    locale,
                    hierarchy_level=(
                        statement.get("level", 1)
                        if key == "workstation_hierarchy" and isinstance(statement, dict)
                        else None
                    ),
                )

    if not rendered_sections:
        raise DocumentRenderingError(f"No reviewable {kind} content was provided.")

    if references:
        # A separate break-only paragraph can overflow to a new page and then
        # break again, leaving a blank page. Attach the break to the heading.
        heading = _insert_paragraph_before(marker, _label(locale, "source_references"), "Heading 1")
        heading.paragraph_format.page_break_before = True
        rows = [
            (
                reference["label"],
                "\n".join(
                    value
                    for value in [
                        reference["source"],
                        (
                            reference["location"]
                            if str(reference["location"]).casefold()
                            not in {"", "document"}
                            else ""
                        ),
                    ]
                    if value
                ),
                reference["quote"],
            )
            for reference in references.values()
        ]
        _insert_table_before(
            document,
            marker,
            [
                _label(locale, "reference"),
                _label(locale, "source"),
                _label(locale, "excerpt"),
            ],
            rows,
            [Inches(1.2), Inches(2.1), Inches(3.45)],
        )


SYSTEM_TEMPLATE_TITLES = {
    "offer": {"en": "Proposal", "de": "Angebot", "fr": "Offre"},
    "requirements_specification": {
        "en": "Technical Specification",
        "de": "Technische Spezifikation",
        "fr": "Spécification technique",
    },
}


def _default_template_bytes(kind, locale="en"):
    document = DocxDocument()
    locale = locale if locale in {"en", "de", "fr"} else "en"
    title = SYSTEM_TEMPLATE_TITLES[kind][locale]
    document.add_heading(title, 0)
    document.add_paragraph(CONTENT_MARKER)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def system_template_bytes(kind, locale="en"):
    if kind not in {"offer", "requirements_specification"}:
        raise DocumentRenderingError("Unsupported project document kind.")
    return _default_template_bytes(kind, locale)


def _template_bytes(template, kind, locale="en"):
    if not template:
        return _default_template_bytes(kind, locale)
    template.template_file.open("rb")
    try:
        return template.template_file.read()
    finally:
        template.template_file.close()


def product_requirement_statements(snapshot, locale):
    """Project reviewed attribute values exactly; do not ask AI to restate them."""
    statements = []
    for row in (snapshot or {}).get("customer_demands", []):
        fields = row.get("custom_fields") or {}
        definitions = fields.get("_reviewed_attribute_definitions") or {}
        for key, definition in definitions.items():
            value = fields.get(key)
            if value in (None, ""):
                continue
            label_field = f"label_{locale}" if definition.get(f"label_{locale}") else "label"
            label = definition.get(label_field) or key
            source_label = label
            if definition.get("unit"):
                label = f"{label} ({definition['unit']})"
            display = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
            evidence = {"source_type": "customer_demand", "source_id": str(row["id"])}
            statements.append({
                "text": f"{row['name']} — {label}: {display}",
                "unresolved": False,
                "evidence": [
                    {**evidence, "location": "name", "quote": json.dumps(row["name"], ensure_ascii=False)},
                    {**evidence, "location": f"custom_fields.{key}", "quote": json.dumps(value, ensure_ascii=False)},
                    {**evidence, "location": f"custom_fields._reviewed_attribute_definitions.{key}.{label_field}",
                     "quote": json.dumps(source_label, ensure_ascii=False)},
                ] + ([{
                    **evidence, "location": f"custom_fields._reviewed_attribute_definitions.{key}.unit",
                    "quote": json.dumps(definition["unit"], ensure_ascii=False),
                }] if definition.get("unit") else []),
            })
    return statements


def order_workstation_hierarchy(rows):
    """Present each occurrence tree parent-first without changing its topology."""
    if not rows or any(
        not row.get("occurrence_id") or not row.get("workstation_id") for row in rows
    ):
        return rows

    def identity(row):
        return str(row["workstation_id"]), str(row["occurrence_id"])
    indexed = {identity(row): row for row in rows}
    if len(indexed) != len(rows):
        return rows  # Validation, not presentation, owns malformed topologies.
    station_order = {station: index for index, station in enumerate(dict.fromkeys(
        str(row["workstation_id"]) for row in rows
    ))}
    children, roots = {}, []
    for row in rows:
        parent = row.get("parent_occurrence_id")
        if parent is None:
            roots.append(row)
        else:
            key = (str(row["workstation_id"]), str(parent))
            if key not in indexed:
                return rows
            children.setdefault(key, []).append(row)
    ordered, visited = [], set()
    stack = list(reversed(sorted(roots, key=lambda row: (
        station_order[str(row["workstation_id"])], row.get("sort_order", 0)
    ))))
    while stack:
        row = stack.pop()
        key = identity(row)
        if key in visited:
            return rows
        visited.add(key)
        ordered.append(row)
        stack.extend(reversed(sorted(
            children.get(key, []), key=lambda child: child.get("sort_order", 0)
        )))
    return ordered if len(ordered) == len(rows) else rows


def build_docx_bytes(
    kind,
    content,
    template=None,
    locale="en",
    source_snapshot=None,
    revision=None,
):
    if "workstation_hierarchy" in content:
        content = {
            **content,
            "workstation_hierarchy": order_workstation_hierarchy(content["workstation_hierarchy"]),
        }
    document = DocxDocument(io.BytesIO(_template_bytes(template, kind, locale)))
    marker = next(
        (paragraph for paragraph in document.paragraphs if CONTENT_MARKER in paragraph.text),
        None,
    )
    if not marker:
        raise DocumentRenderingError("The selected template has no content marker.")

    _insert_document_information(document, marker, source_snapshot, locale, revision)
    if isinstance(content.get("sections"), list):
        for section in content["sections"]:
            title = section.get("title") or section.get("name")
            if title:
                _insert_paragraph_before(marker, title, "Heading 1")
            _render_generic_value(
                marker,
                section.get("content", section.get("body", section)),
            )
    else:
        _render_structured_content(
            document,
            marker,
            kind,
            content,
            locale,
            source_snapshot=source_snapshot,
        )
    marker._element.getparent().remove(marker._element)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def convert_docx_to_pdf(docx_bytes):
    executable = settings.PROJECT_CONTENT_LIBREOFFICE_BINARY
    resolved = shutil.which(executable) or (executable if os.path.exists(executable) else None)
    if not resolved:
        raise DocumentRenderingError("LibreOffice is not available for PDF rendering.")

    with tempfile.TemporaryDirectory(prefix="project-content-") as directory:
        docx_path = os.path.join(directory, "source.docx")
        with open(docx_path, "wb") as handle:
            handle.write(docx_bytes)
        process = subprocess.run(
            [
                resolved,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                directory,
                docx_path,
            ],
            capture_output=True,
            timeout=settings.PROJECT_CONTENT_RENDER_TIMEOUT_SECONDS,
            check=False,
        )
        pdf_path = os.path.join(directory, "source.pdf")
        if process.returncode or not os.path.exists(pdf_path):
            raise DocumentRenderingError(
                "LibreOffice could not render the generated DOCX."
            )
        with open(pdf_path, "rb") as handle:
            return handle.read()


def _safe_project_code(project):
    candidate = "".join(
        char if char.isalnum() or char in {"-", "_"} else "-"
        for char in project.project_name
    ).strip("-")
    return candidate[:80] or "Project"


def generated_file_names(project, kind, revision):
    label = "Proposal" if kind == "offer" else "Technical-Specification"
    base = f"{_safe_project_code(project)}-{label}-R{revision}-{date.today():%Y%m%d}"
    return f"{base}.docx", f"{base}.pdf"


@transaction.atomic
def render_generation_run(run, user, content=None):
    run = ProjectDocumentGenerationRun.objects.select_for_update().select_related(
        "project",
        "generated_document",
    ).get(pk=run.pk)
    content = content or run.draft_content or run.result
    if "workstation_hierarchy" in content:
        content = {
            **content,
            "workstation_hierarchy": order_workstation_hierarchy(content["workstation_hierarchy"]),
        }
    template = (
        run.project.company_profile.project_document_templates.filter(
            kind=run.kind,
            locale=run.locale,
            is_active=True,
            validation_status="valid",
        ).first()
        if run.project.company_profile_id
        else None
    )
    template_version = document_template_version(template)
    existing = GeneratedProjectDocumentRevision.objects.filter(
        generated_document=run.generated_document,
        source_hash=run.source_hash,
        template_version=template_version,
    ).first()
    if existing:
        run.applied_revision = existing
        run.status = ProjectDocumentGenerationRun.Status.APPLIED
        run.save(update_fields=["applied_revision", "status", "updated_at"])
        return existing

    next_revision = (
        run.generated_document.revisions.aggregate(value=Max("revision"))["value"] or 0
    ) + 1
    docx_bytes = build_docx_bytes(
        run.kind,
        content,
        template,
        run.locale,
        source_snapshot=run.source_snapshot,
        revision=next_revision,
    )
    pdf_bytes = convert_docx_to_pdf(docx_bytes)
    folder = get_project_content_folder(run.project, run.kind, run.locale)
    docx_name, pdf_name = generated_file_names(run.project, run.kind, next_revision)
    stored_files = []
    try:
        docx_document = Document(
            folder=folder,
            author=user,
            assignee=user,
            created_by=user,
            file_name=docx_name,
            document_type=(
                "Offer/quotation:OFF"
                if run.kind == "offer"
                else "General and detailed specifications:SPC"
            ),
            status="Draft",
            revision=next_revision,
        )
        docx_document.file.save(docx_name, ContentFile(docx_bytes), save=False)
        stored_files.append((docx_document.file.storage, docx_document.file.name))
        docx_document.save()

        pdf_document = Document(
            folder=folder,
            author=user,
            assignee=user,
            created_by=user,
            file_name=pdf_name,
            document_type=docx_document.document_type,
            status="Draft",
            revision=next_revision,
        )
        pdf_document.file.save(pdf_name, ContentFile(pdf_bytes), save=False)
        stored_files.append((pdf_document.file.storage, pdf_document.file.name))
        pdf_document.save()

        revision = GeneratedProjectDocumentRevision.objects.create(
            generated_document=run.generated_document,
            revision=next_revision,
            source_hash=run.source_hash,
            template=template,
            template_version=template_version,
            structured_content=content,
            validation_report={"valid": True, "warnings": run.warnings},
            docx_document=docx_document,
            pdf_document=pdf_document,
            docx_hash=hashlib.sha256(docx_bytes).hexdigest(),
            pdf_hash=hashlib.sha256(pdf_bytes).hexdigest(),
            created_by=user,
        )
        run.generated_document.current_revision = revision
        run.generated_document.status = run.generated_document.Status.DRAFT
        run.generated_document.source_hash = run.source_hash
        run.generated_document.save(
            update_fields=["current_revision", "status", "source_hash", "updated_at"]
        )
        run.applied_revision = revision
        run.status = ProjectDocumentGenerationRun.Status.APPLIED
        run.draft_content = content
        run.save(
            update_fields=[
                "applied_revision",
                "status",
                "draft_content",
                "updated_at",
            ]
        )
        record_project_content_event(
            project=run.project,
            event_type="document_revision_created",
            instance=revision,
            actor=user,
            metadata={
                "kind": run.kind,
                "revision": revision.revision,
                "template_version": template_version,
            },
        )
        return revision
    except Exception:
        for storage, name in stored_files:
            try:
                storage.delete(name)
            except Exception:
                pass
        raise
