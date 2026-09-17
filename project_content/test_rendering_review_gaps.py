from io import BytesIO

from django.test import SimpleTestCase
from docx import Document

from project_content.rendering import build_docx_bytes


class RenderingReviewGapTests(SimpleTestCase):
    def test_references_break_belongs_to_heading_not_empty_paragraph(self):
        content = {"executive_summary": [{"text": "A product.", "evidence": [
            {"source_type": "document", "source_id": "1", "quote": "A product."}
        ]}]}
        document = Document(BytesIO(build_docx_bytes("offer", content)))
        heading = next(p for p in document.paragraphs if p.text == "Source References")
        self.assertTrue(heading.paragraph_format.page_break_before)
        self.assertFalse(document.element.xpath('.//w:br[@w:type="page"]'))

    def test_open_prerequisites_are_visible_next_to_acceptance_conditions(self):
        content = {
            "acceptance_conditions": [{"text": "Verify torque within 10 percent.",
                                       "unresolved": False}],
            "unresolved_decisions": ["Torque measurement procedure is unspecified."],
        }
        document = Document(BytesIO(build_docx_bytes("requirements_specification", content)))
        paragraphs = [p.text for p in document.paragraphs]
        start = paragraphs.index("Acceptance Conditions")
        criterion = paragraphs.index("Verify torque within 10 percent.")
        self.assertIn("Unresolved: Torque measurement procedure is unspecified.",
                      paragraphs[start:criterion])
        self.assertFalse(content["acceptance_conditions"][0]["unresolved"])

    def test_no_prerequisites_are_invented_for_complete_content(self):
        document = Document(BytesIO(build_docx_bytes("requirements_specification", {
            "acceptance_conditions": [{"text": "Verify length."}],
        })))
        self.assertNotIn("Unresolved:", "\n".join(p.text for p in document.paragraphs))
