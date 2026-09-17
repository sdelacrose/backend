from datetime import date
import json
import shutil
import tempfile
import zipfile
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.core.management import call_command
from django.db.models.signals import post_save
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from docx import Document as DocxDocument

from companyprofileapis.models import Opportunity
from models.models import CompanyProfile, User
from project_management.models import (
    Document,
    DocumentContentSegment,
    DocumentSearchIndex,
    ECPMPConfig,
    Folder,
    Project,
    Requirement,
    ProjectWorkstation,
    CompanySolutionSpaceItem,
    ProjectSolutionSpaceItem,
    ProcessAssemblyExtractionRun,
    ProjectWorkstationNode,
    WorkstationRevision,
    WorkstationSnapshot,
)
from project_management.serializers.document import FolderSerializer
from project_management.signals import create_default_folders

from .folders import ensure_project_content_folders
from .models import (
    AcceptanceChecklist,
    AcceptanceCriterion,
    AcceptanceGenerationRun,
    GeneratedProjectDocument,
    GeneratedProjectDocumentRevision,
    ProjectContentAuditEvent,
    ProjectDocumentGenerationRun,
)
from .rendering import build_docx_bytes, document_template_version
from .serializers import (
    AcceptanceCriterionCommentSerializer,
    AcceptanceCriterionSerializer,
)
from .services import (
    ProjectContentError,
    apply_acceptance_generation_run,
    apply_document_generation_run,
    approve_document_revision,
    create_acceptance_generation_run,
    create_document_generation_run,
    process_acceptance_generation_run,
    process_document_generation_run,
    queue_document_generation_run,
    _normalize_result_references,
    _validate_requirements_specification_coverage,
    _validate_result_references,
    _attachment_manifest_snapshot,
    _validate_offer_attachment_manifest,
    _validate_requirements_specification_context,
)
from .snapshots import _tree_rows, build_project_content_snapshot, canonical_hash
from .template_validation import CONTENT_MARKER, TemplateValidationError, validate_template_file
from .tasks import _mark_generation_retry_pending


TEST_STORAGES = {
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
}


@override_settings(
    PROJECT_DOCUMENT_GENERATION_ENABLED=True,
    PROJECT_DOCUMENT_AI_ENABLED=True,
    ACCEPTANCE_CHECKLISTS_ENABLED=True,
    PRODUCT_ASSEMBLY_EXTRACTION_SERVICE_KEY="test-service-key",
    STORAGES=TEST_STORAGES,
)
class ProjectContentTests(TestCase):
    @override_settings(PROJECT_CONTENT_RUN_SYNCHRONOUSLY=True)
    @patch("ecpmpapp.permissions.LicensePermission.has_permission", return_value=True)
    @patch("project_content.views.require_project_content_access")
    @patch("project_content.services.requests.post")
    def test_document_generation_uses_only_the_typed_structured_evidence_identity(
        self, post, _access, _license
    ):
        client = APIClient()
        client.force_authenticate(self.user)
        result = self.offer_result()
        statement = result["content"]["executive_summary"][0]
        statement["text"] = self.project.project_name
        statement["evidence"] = [{
            "source_type": "project", "source_id": str(self.project.id),
            "document_id": "", "segment_id": "", "requirement_id": "",
            "quote": self.project.project_name,
        }]
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        response = client.post(
            f"/api/project-content/projects/{self.project.id}/offer/generate/",
            {"locale": "en"}, format="json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        self.assertEqual(response.data["status"], "review_ready", response.content)

        for locale, field, error_code in (
            ("de", "source_id", "unknown_structured_reference"),
            ("fr", "document_id", "unknown_document_reference"),
        ):
            with self.subTest(field=field):
                statement["evidence"][0] = {
                    "source_type": "project", "source_id": str(self.project.id),
                    "document_id": None, "segment_id": None, "requirement_id": None,
                    "quote": self.project.project_name, field: "not-in-project",
                }
                rejected = client.post(
                    f"/api/project-content/projects/{self.project.id}/offer/generate/",
                    {"locale": locale}, format="json",
                )
                self.assertEqual(rejected.data["status"], "failed", rejected.content)
                self.assertEqual(rejected.data["error_code"], error_code, rejected.content)

    @patch("project_content.services.requests.post")
    def test_specification_generates_from_existing_process_tree_without_duplicate_workstations(self, post):
        from project_management.models import (
            SolutionDemand, ProductAttributeExtractionRun,
            ProductAttributeAppliedColumn, ProductAttributeAppliedValue,
        )
        demand = SolutionDemand.objects.create(
            project=self.project, subject="Precision", custom_fields={"payload": "10 kg"},
        )
        attribute_run = ProductAttributeExtractionRun.objects.create(
            project=self.project, status="applied", source_snapshot_hash="source",
            matrix_snapshot_hash="matrix", created_by=self.user, modified_by=self.user,
        )
        column = ProductAttributeAppliedColumn.objects.create(
            run=attribute_run, proposal_key="payload", column_key="payload",
            applied_payload={"label": "Rated payload", "label_de": "Nutzlast", "unit": "kg"},
        )
        ProductAttributeAppliedValue.objects.create(
            applied_column=column, solution_demand=demand, value="15 kg",
        )
        foreign_project = Project.objects.create(project_name="Other project", contract_value=0)
        foreign_demand = SolutionDemand.objects.create(
            project=foreign_project, subject="Other product", custom_fields={"payload": "999 kg"},
        )
        ProductAttributeAppliedValue.objects.create(
            applied_column=column, solution_demand=foreign_demand, value="999 kg",
        )
        run, _ = create_document_generation_run(
            project=self.project, user=self.user, kind="requirements_specification",
            document_ids=[self.document.id],
        )
        self.assertFalse(run.source_snapshot["project_workstations"])
        station = run.source_snapshot["process_stations"][0]
        evidence = {"source_type": "process_station", "source_id": str(station["id"]), "quote": station["name"]}
        statement = {"text": station["name"], "evidence": [evidence]}
        result = {"content": {
            "purpose_and_scope": [statement], "functional_requirements": [statement],
            "selected_workstations": [statement], "evidence_metadata": [statement],
            "workstation_hierarchy": [{**statement, "workstation_id": str(station["id"]),
                "occurrence_id": str(station["id"]), "parent_occurrence_id": None,
                "level": 1, "sort_order": 0}],
            "traceability_matrix": [{"requirement_id": str(self.requirement.id),
                "coverage": station["name"], "evidence": [evidence]}],
            "omitted_requirement_ids": [], "assumptions": [], "exclusions": [],
        }}
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        processed = process_document_generation_run(run.id)
        self.assertEqual(processed.status, "review_ready", processed.error_message)
        attributes = processed.draft_content.get("product_requirements", [])
        self.assertEqual([row["text"] for row in attributes], ["Precision — Rated payload (kg): 10 kg"])
        from .rendering import product_requirement_statements
        self.assertEqual(
            product_requirement_statements(run.source_snapshot, "de")[0]["text"],
            "Precision — Nutzlast (kg): 10 kg",
        )
        self.assertNotIn("999 kg", json.dumps(run.source_snapshot))
        rendered = DocxDocument(BytesIO(build_docx_bytes(
            "requirements_specification", processed.draft_content,
            source_snapshot=run.source_snapshot,
        )))
        text = "\n".join(p.text for p in rendered.paragraphs)
        self.assertIn("Product Requirements", text)
        self.assertIn("Precision — Rated payload (kg): 10 kg", text)
        self.assertFalse(ProjectWorkstation.objects.filter(project=self.project).exists())
        invalid_run, _ = create_document_generation_run(
            project=self.project, user=self.user, kind="requirements_specification",
            document_ids=[self.document.id], locale="de",
        )
        result["content"]["workstation_hierarchy"][0]["parent_occurrence_id"] = str(station["id"])
        rejected = process_document_generation_run(invalid_run.id)
        self.assertEqual(rejected.status, "failed")
        self.assertEqual(rejected.error_code, "invalid_workstation_hierarchy")

    @override_settings(PROJECT_DOCUMENT_AI_TIMEOUT_SECONDS=315, PROJECT_CONTENT_AI_TIMEOUT_SECONDS=137)
    @patch("project_content.services.requests.post")
    def test_document_failures_distinguish_timeout_from_invalid_evidence(self, post):
        for http_status, detail, code, text in (
            (504, "Timed out", "generation_service_timeout", "timed out"),
            (422, {"code": "generation_validation_failed"},
             "generation_validation_failed", "project evidence"),
        ):
            with self.subTest(http_status=http_status):
                post.return_value = SimpleNamespace(
                    status_code=http_status, json=lambda: {"detail": detail},
                )
                run, _ = create_document_generation_run(
                    project=self.project, user=self.user, kind="offer",
                    document_ids=[self.document.id],
                )
                result = process_document_generation_run(run.id)
                self.assertEqual(result.status, "failed")
                self.assertEqual(result.error_code, code)
                self.assertIn(text, result.error_message)
                self.assertNotIn("acceptance tests", result.error_message)
                self.assertEqual(post.call_args.kwargs["timeout"], 315)

    def setUp(self):
        post_save.disconnect(create_default_folders, sender=Project)
        self.addCleanup(post_save.connect, create_default_folders, sender=Project)
        self.notification_patcher = patch(
            "ecpmpapp.utils.should_notify_user_comprehensive",
            return_value=False,
        )
        self.notification_patcher.start()
        self.addCleanup(self.notification_patcher.stop)
        self.matrix_room_patcher = patch(
            "chat.signals.matrix_service.ensure_room",
            return_value=SimpleNamespace(room_type="project"),
        )
        self.matrix_room_patcher.start()
        self.addCleanup(self.matrix_room_patcher.stop)
        self.media_root = tempfile.mkdtemp(prefix="project-content-tests-")
        self.media_settings = self.settings(MEDIA_ROOT=self.media_root)
        self.media_settings.enable()
        self.addCleanup(self.media_settings.disable)
        self.addCleanup(shutil.rmtree, self.media_root, True)

        self.user = User.objects.create_user(
            email="project-content-owner@example.com",
            password="password",
            first_name="Project",
            last_name="Owner",
            user_type="company_user",
        )
        self.company = CompanyProfile.objects.create(
            admin_user=self.user,
            company_name="Project Content Company",
            company_address="Street 1",
            company_town="Town",
            company_tell_no="123",
            company_mob_no="456",
            company_web_address="https://example.com",
            company_email_address=self.user.email,
            company_tax_reference="tax",
            company_iban="iban",
        )
        self.user.company = self.company
        self.user.save(update_fields=["company"])
        self.project = Project.objects.create(
            project_name="DM-200",
            project_description="Bench drilling machine project",
            contract_value=1000,
            company_profile=self.company,
            user=self.user,
            created_by=self.user,
        )
        self.root = Folder.objects.create(
            project=self.project,
            name=f"{self.project.id}_root_folder",
        )
        self.source_folder = Folder.objects.create(
            project=self.project,
            parent=self.root,
            name="Customer Demands",
        )
        self.document = Document.objects.create(
            folder=self.source_folder,
            author=self.user,
            created_by=self.user,
            file_name="DM-200 Demand.docx",
            document_type="Customer Demand:CDE",
            status="Draft",
        )
        self.index = DocumentSearchIndex.objects.create(
            document=self.document,
            status=DocumentSearchIndex.Status.SEARCHABLE,
            content_hash="project-content-hash-v1",
            parser_version="document-search-v1",
        )
        self.segment = DocumentContentSegment.objects.create(
            document=self.document,
            index=self.index,
            segment_type=DocumentContentSegment.SegmentType.PAGE,
            sequence=1,
            location_label="Page 1",
            text="REQ-001 requires a 400 V supply and a witnessed safety test.",
        )
        self.requirement = Requirement.objects.create(
            project=self.project,
            name="Supply voltage",
            object_number="REQ-001",
            custom_fields={"value": "400 V"},
            created_by=self.user,
            modified_by=self.user,
        )

    def test_offer_approval_requires_a_current_approved_specification_attachment(self):
        generated = GeneratedProjectDocument.objects.create(project=self.project, kind="offer")
        revision = GeneratedProjectDocumentRevision.objects.create(
            generated_document=generated, revision=1, source_hash="draft",
            template_version="system-v1", structured_content={"attachment_manifest": []},
            docx_document=self.document, pdf_document=self.document,
            docx_hash="docx", pdf_hash="pdf", created_by=self.user,
        )
        with self.assertRaises(ProjectContentError) as context:
            approve_document_revision(revision, self.user)
        self.assertEqual(context.exception.code, "technical_specification_required")

    def test_offer_approval_rejects_an_unapproved_or_stale_specification(self):
        specification = self.approved_specification_fixture()
        generated = GeneratedProjectDocument.objects.create(project=self.project, kind="offer")
        revision = GeneratedProjectDocumentRevision.objects.create(
            generated_document=generated, revision=1, source_hash="draft",
            template_version="system-v1",
            structured_content={"attachment_manifest": self.specification_manifest},
            docx_document=self.document, pdf_document=self.document,
            docx_hash="docx", pdf_hash="pdf", created_by=self.user,
        )
        for approval, state in [("draft", "current"), ("approved", "out_of_date")]:
            with self.subTest(approval=approval, state=state):
                specification.approval_status = approval
                specification.save(update_fields=["approval_status"])
                specification.generated_document.status = state
                specification.generated_document.save(update_fields=["status"])
                with self.assertRaises(ProjectContentError) as context:
                    approve_document_revision(revision, self.user)
                self.assertEqual(context.exception.code, "technical_specification_required")

    def approved_specification_fixture(self):
        generated = GeneratedProjectDocument.objects.create(
            project=self.project, kind="requirements_specification", status="current",
        )
        pdf = Document.objects.create(
            folder=self.source_folder, author=self.user, created_by=self.user,
            file_name="Technical-Specification-R1.pdf",
        )
        revision = GeneratedProjectDocumentRevision.objects.create(
            generated_document=generated, revision=1, source_hash="specification",
            template_version="system-v1", structured_content={},
            docx_document=pdf, pdf_document=pdf, docx_hash="docx", pdf_hash="pdf",
            created_by=self.user, approval_status="approved",
        )
        generated.current_revision = revision
        generated.save(update_fields=["current_revision"])
        self.specification_manifest = [{
            "id": str(revision.id), "name": pdf.file_name,
            "document_type": "technical_specification", "version": 1,
            "generation_status": "current",
        }]
        return revision

    def test_specification_approval_rejects_unresolved_station_coverage(self):
        revision = self.approved_specification_fixture()
        revision.approval_status = "draft"
        revision.save(update_fields=["approval_status"])
        ProcessAssemblyExtractionRun.objects.create(
            project=self.project, status="applied", input_snapshot_hash="source",
            result={"rows": [], "unresolved_items": ["No selected template for WS30."]},
        )
        with self.assertRaises(ProjectContentError) as context:
            approve_document_revision(revision, self.user)
        self.assertEqual(context.exception.code, "incomplete_station_coverage")

    def test_acceptance_generation_rejects_unresolved_station_coverage(self):
        ProcessAssemblyExtractionRun.objects.create(
            project=self.project, status="applied", input_snapshot_hash="source",
            result={"rows": [], "unresolved_items": ["No selected template for WS30."]},
        )
        with self.assertRaises(ProjectContentError) as context:
            create_acceptance_generation_run(project=self.project, user=self.user)
        self.assertEqual(context.exception.code, "incomplete_station_coverage")

    def test_reviewed_optional_scope_note_does_not_block_specification_approval(self):
        from project_management.services.process_requirements_extraction import confirm_station_coverage

        revision = self.approved_specification_fixture()
        run = ProcessAssemblyExtractionRun.objects.create(
            project=self.project, status="applied", input_snapshot_hash="source",
            result={"rows": [{"name": self.requirement.name}],
                "unresolved_items": ["An optional future marking cell is excluded."]},
        )
        with self.assertRaises(ProjectContentError):
            approve_document_revision(revision, self.user)
        confirm_station_coverage(run, self.user, "All base stations reviewed; optional cell is outside this project.")
        self.assertEqual(approve_document_revision(revision, self.user).approval_status, "approved")
        self.requirement.name = "Changed scope"
        self.requirement.save(update_fields=["name"])
        with self.assertRaises(ProjectContentError) as context:
            approve_document_revision(revision, self.user)
        self.assertEqual(context.exception.code, "incomplete_station_coverage")
        confirm_station_coverage(run, self.user, "Rechecked, but the original required station is still absent.")
        with self.assertRaises(ProjectContentError):
            approve_document_revision(revision, self.user)


    def test_a_fresh_review_does_not_clear_coverage_until_its_stations_exist(self):
        ProcessAssemblyExtractionRun.objects.create(
            project=self.project, status="review_ready", input_snapshot_hash="source",
            result={"rows": [{"name": "WS30 Saw", "selected": True}], "unresolved_items": []},
        )
        with self.assertRaises(ProjectContentError) as context:
            create_acceptance_generation_run(project=self.project, user=self.user)
        self.assertEqual(context.exception.code, "incomplete_station_coverage")

    def test_specification_approval_requires_a_workstation_for_a_linked_leaf_requirement(self):
        template = CompanySolutionSpaceItem.objects.create(
            company_profile=self.company, assembly_name="Storage",
            assembly_type="production_lines",
        )
        selected = ProjectSolutionSpaceItem.objects.create(
            project=self.project, assembly_name="Storage", assembly_type="production_lines",
            source_company_item=template,
        )
        selected.requirements.add(self.requirement)
        revision = self.approved_specification_fixture()
        with self.assertRaises(ProjectContentError) as context:
            approve_document_revision(revision, self.user)
        self.assertEqual(context.exception.code, "missing_project_workstations")

    def test_specification_approval_requires_the_current_workstation_bom(self):
        from project_management.services.workstation_composition import create_project_workstation
        from project_management.services.workstation_bom import create_bom_plan, apply_bom_plan
        template = CompanySolutionSpaceItem.objects.create(
            company_profile=self.company, assembly_name="Storage", assembly_type="production_lines",
        )
        workstation = create_project_workstation(
            project=self.project, source=template, actor=self.user, idempotency_key="qa-storage",
        ).workstation
        revision = self.approved_specification_fixture()
        with self.assertRaises(ProjectContentError) as context:
            approve_document_revision(revision, self.user)
        self.assertEqual(context.exception.code, "incomplete_workstation_bom")
        plan = create_bom_plan(
            project=self.project, workstation_ids=[workstation.id], actor=self.user,
            idempotency_key="qa-storage-bom",
        )
        apply_bom_plan(plan=plan, actor=self.user, expected_target_hash=plan.target_hash)
        approved = approve_document_revision(revision, self.user)
        self.assertEqual(approved.approval_status, "approved")

    def test_new_workstation_marks_the_approved_specification_out_of_date(self):
        from project_management.services.workstation_composition import create_project_workstation
        revision = self.approved_specification_fixture()
        template = CompanySolutionSpaceItem.objects.create(
            company_profile=self.company, assembly_name="Storage", assembly_type="production_lines",
        )
        create_project_workstation(
            project=self.project, source=template, actor=self.user, idempotency_key="new-station",
        )
        revision.generated_document.refresh_from_db()
        self.assertEqual(revision.generated_document.status, "out_of_date")

    def offer_result(self):
        evidence = {
            "document_id": str(self.document.id),
            "segment_id": str(self.segment.id),
            "source_type": "document",
            "source_id": None,
            "requirement_id": None,
            "location": "Page 1",
            "quote": "REQ-001 requires a 400 V supply",
        }
        return {
            "schema_version": "1",
            "project_language": "en",
            "content": {
                "attachment_manifest": getattr(self, "specification_manifest", []),
                "executive_summary": [
                    {
                        "text": "A 400 V supply is required.",
                        "evidence": [evidence],
                        "unresolved": False,
                    }
                ],
                "open_questions": ["Commercial price is not available."],
            },
            "warnings": ["Commercial values require review."],
        }

    def acceptance_result(self, action="add"):
        existing = AcceptanceCriterion.objects.filter(
            checklist__project=self.project,
            stable_key="safety-test",
        ).first()
        return {
            "schema_version": "1",
            "project_language": "en",
            "criteria": [
                {
                    "proposal_key": "safety-test",
                    "change_action": action,
                    "existing_criterion_id": str(existing.id) if existing else None,
                    "phases": ["both"],
                    "sort_order": 0,
                    "name": "Witnessed safety test",
                    "acceptance_criteria": "The safety test is witnessed and passed.",
                    "linked_requirement_ids": [str(self.requirement.id)],
                    "priority": None,
                    "evidence": [
                        {
                            "document_id": str(self.document.id),
                            "segment_id": str(self.segment.id),
                            "source_type": "document",
                            "source_id": None,
                            "requirement_id": str(self.requirement.id),
                            "location": "Page 1",
                            "quote": "a witnessed safety test",
                        }
                    ],
                    "warnings": [],
                }
            ],
            "warnings": [],
            "unresolved_items": [],
        }

    def test_snapshot_is_json_serializable_and_deterministic(self):
        first, first_hash = build_project_content_snapshot(self.project)
        second, second_hash = build_project_content_snapshot(self.project)
        json.dumps(first)
        self.assertEqual(first, second)
        self.assertEqual(first_hash, second_hash)
        self.assertEqual(first["requirements"][0]["id"], str(self.requirement.id))

    def test_snapshot_normalizes_zero_contract_value_as_missing(self):
        self.project.contract_value = 0
        self.project.save(update_fields=["contract_value"])
        snapshot, _ = build_project_content_snapshot(self.project)
        self.assertIsNone(snapshot["project"]["contract_value"])

    def test_snapshot_contains_selected_workstation_revision_and_hierarchy(self):
        workstation = ProjectWorkstation.objects.create(
            project=self.project,
            name="Assembly station 1",
            template_hash="template-hash",
            created_by=self.user,
        )
        revision = WorkstationRevision.objects.create(
            workstation=workstation,
            number=1,
            source_hash="source-hash",
            reason="Initial reviewed workstation",
            created_by=self.user,
        )
        WorkstationSnapshot.objects.create(
            revision=revision,
            content={"version": 1},
            content_hash="snapshot-hash",
        )
        root = ProjectWorkstationNode.objects.create(
            workstation=workstation,
            revision=revision,
            name="Assembly 1.1",
            level=1,
            sort_order=0,
            quantity=2,
            unit="pcs",
            created_by=self.user,
        )
        child = ProjectWorkstationNode.objects.create(
            workstation=workstation,
            revision=revision,
            parent=root,
            name="Sub-assembly 1.1.1",
            level=2,
            sort_order=1,
            quantity=3,
            unit="pcs",
            role="added",
            created_by=self.user,
        )

        snapshot, _ = build_project_content_snapshot(self.project)

        self.assertEqual(snapshot["project_workstations"][0]["active_revision"], 1)
        self.assertEqual(snapshot["project_workstations"][0]["snapshot_hash"], "snapshot-hash")
        self.assertEqual(snapshot["project_workstations"][0]["nodes"][1]["parent_occurrence_id"], str(root.occurrence_id))
        self.assertEqual(snapshot["project_workstations"][0]["nodes"][1]["occurrence_id"], str(child.occurrence_id))

    def test_offer_attachment_manifest_is_exactly_the_current_technical_specification(self):
        generated = GeneratedProjectDocument.objects.create(
            project=self.project,
            kind="requirements_specification",
            status="current",
        )
        docx = Document.objects.create(
            folder=self.source_folder,
            author=self.user,
            created_by=self.user,
            file_name="DM-200-Technical-Specification-R2.docx",
            document_type="General and detailed specifications:SPC",
            revision=2,
        )
        pdf = Document.objects.create(
            folder=self.source_folder,
            author=self.user,
            created_by=self.user,
            file_name="DM-200-Technical-Specification-R2.pdf",
            document_type="General and detailed specifications:SPC",
            revision=2,
        )
        revision = GeneratedProjectDocumentRevision.objects.create(
            generated_document=generated,
            revision=2,
            source_hash="technical-source-hash",
            template_version="system-v1",
            structured_content={"purpose_and_scope": []},
            docx_document=docx,
            pdf_document=pdf,
            docx_hash="docx-hash",
            pdf_hash="pdf-hash",
            created_by=self.user,
        )
        generated.current_revision = revision
        generated.status = "current"
        generated.save(update_fields=["current_revision", "status"])

        manifest = _attachment_manifest_snapshot(self.project)
        self.assertEqual(
            manifest,
            [{
                "id": str(revision.id),
                "name": "DM-200-Technical-Specification-R2.pdf",
                "document_type": "technical_specification",
                "version": 2,
                "generation_status": "current",
            }],
        )
        _validate_offer_attachment_manifest(
            {"attachment_manifest": manifest},
            {"content": {"attachment_manifest": manifest}},
        )
        with self.assertRaises(ProjectContentError):
            _validate_offer_attachment_manifest(
                {"attachment_manifest": manifest},
                {"content": {"attachment_manifest": []}},
            )

    def test_technical_specification_fixed_headings_are_localized(self):
        content = {
            "purpose_and_scope": ["Scope text"],
            "selected_workstations": ["Station 1"],
            "unresolved_decisions": ["Decision 1"],
        }
        expectations = {
            "en": {"Technical Specification", "Purpose and Scope", "Selected Workstations", "Unresolved Decisions"},
            "de": {"Technische Spezifikation", "Zweck und Umfang", "Ausgewählte Arbeitsstationen", "Offene Entscheidungen"},
            "fr": {"Spécification technique", "Objectif et périmètre", "Postes de travail sélectionnés", "Décisions en suspens"},
        }
        for locale, expected in expectations.items():
            rendered = build_docx_bytes("requirements_specification", content, locale=locale)
            document = DocxDocument(BytesIO(rendered))
            text = {paragraph.text for paragraph in document.paragraphs}
            self.assertTrue(expected.issubset(text), (locale, text))

    def test_technical_specification_requires_complete_workstation_sections(self):
        snapshot = {
            "requirements": [],
            "project_workstations": [
                {
                    "id": "501",
                    "nodes": [
                        {
                            "occurrence_id": "00000000-0000-0000-0000-000000000001",
                            "parent_occurrence_id": None,
                            "name": "Assembly station 1",
                            "level": 1,
                            "sort_order": 0,
                            "is_active": True,
                        },
                        {
                            "occurrence_id": "00000000-0000-0000-0000-000000000002",
                            "parent_occurrence_id": "00000000-0000-0000-0000-000000000001",
                            "name": "Servo module",
                            "level": 2,
                            "sort_order": 1,
                            "is_active": True,
                        },
                    ],
                }
            ],
            "unresolved_decisions": [{"id": "88"}],
        }
        with self.assertRaises(ProjectContentError) as context:
            _validate_requirements_specification_context(
                snapshot,
                {"content": {"traceability_matrix": [], "omitted_requirement_ids": []}},
            )
        self.assertEqual(context.exception.code, "incomplete_technical_specification")

        evidence = {
            "source_type": "project_workstation",
            "source_id": "501",
            "quote": "Assembly station 1 Servo module",
        }
        result = {
            "content": {
                "purpose_and_scope": [{"text": "Scope", "evidence": [evidence]}],
                "selected_workstations": [
                    {"text": "Assembly station 1", "evidence": [evidence]}
                ],
                "workstation_hierarchy": [
                    {
                        "text": "Assembly station 1",
                        "workstation_id": "501",
                        "occurrence_id": "00000000-0000-0000-0000-000000000001",
                        "parent_occurrence_id": None,
                        "level": 1,
                        "sort_order": 0,
                        "evidence": [evidence],
                    },
                    {
                        "text": "Servo module",
                        "workstation_id": "501",
                        "occurrence_id": "00000000-0000-0000-0000-000000000002",
                        "parent_occurrence_id": "00000000-0000-0000-0000-000000000001",
                        "level": 2,
                        "sort_order": 1,
                        "evidence": [evidence],
                    },
                ],
                "functional_requirements": [{"text": "Function", "evidence": [evidence]}],
                "assumptions": [],
                "exclusions": [],
                "evidence_metadata": [{"text": "Revision 1", "evidence": [evidence]}],
                "unresolved_decisions": ["Select one of two matching assemblies."],
                "traceability_matrix": [],
                "omitted_requirement_ids": [],
            }
        }
        _validate_requirements_specification_context(snapshot, result)

        result["content"]["workstation_hierarchy"][1]["parent_occurrence_id"] = None
        with self.assertRaises(ProjectContentError) as context:
            _validate_requirements_specification_context(snapshot, result)
        self.assertEqual(context.exception.code, "invalid_workstation_hierarchy")

        result["content"]["workstation_hierarchy"][1]["parent_occurrence_id"] = (
            "00000000-0000-0000-0000-000000000001"
        )
        result["content"]["workstation_hierarchy"].append(
            dict(result["content"]["workstation_hierarchy"][1])
        )
        with self.assertRaises(ProjectContentError) as context:
            _validate_requirements_specification_context(snapshot, result)
        self.assertEqual(context.exception.code, "invalid_workstation_hierarchy")

    def test_system_folder_backfill_is_idempotent(self):
        first = ensure_project_content_folders(self.project, "de")
        second = ensure_project_content_folders(self.project, "de")
        self.assertTrue(all(row["created"] for row in first))
        self.assertTrue(all(not row["created"] for row in second))
        self.assertEqual(
            set(self.project.folders.exclude(system_key=None).values_list("name", flat=True)),
            {"Angebot", "Lastenheft"},
        )

    @patch("project_content.services.requests.post")
    @patch("project_content.rendering.convert_docx_to_pdf", return_value=b"%PDF-1.7\n")
    def test_offer_generation_review_render_and_approval(self, _pdf, post):
        self.approved_specification_fixture()
        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        run, created = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        self.assertTrue(created)
        process_document_generation_run(run.id)
        run.refresh_from_db()
        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.REVIEW_READY)
        self.assertNotIn(
            "generation_contract",
            post.call_args.kwargs["json"],
        )

        revision = apply_document_generation_run(run, self.user)
        self.assertTrue(revision.docx_document.file_name.endswith(".docx"))
        self.assertTrue(revision.pdf_document.file_name.endswith(".pdf"))
        self.assertEqual(revision.approval_status, revision.ApprovalStatus.DRAFT)
        approve_document_revision(revision, self.user)
        revision.refresh_from_db()
        self.assertEqual(revision.approval_status, revision.ApprovalStatus.APPROVED)
        self.assertTrue(revision.docx_hash)
        self.assertTrue(revision.pdf_hash)

    @patch("project_content.services.requests.post")
    @patch("project_content.rendering.convert_docx_to_pdf", return_value=b"%PDF-1.7\n")
    def test_approved_document_only_becomes_outdated_when_selected_document_changes(self, _pdf, post):
        self.approved_specification_fixture()
        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        run, _ = create_document_generation_run(
            project=self.project, user=self.user, kind="offer",
            document_ids=[self.document.id],
        )
        process_document_generation_run(run.id)
        run.refresh_from_db()
        revision = apply_document_generation_run(run, self.user)
        approve_document_revision(revision, self.user)

        unrelated = Document.objects.create(
            folder=self.source_folder, author=self.user, created_by=self.user,
            file_name="another-generated-output.pdf",
        )
        revision.generated_document.refresh_from_db()
        self.assertEqual(revision.generated_document.status, "current")
        unrelated.file_name = "renamed-output.pdf"
        unrelated.save()
        unrelated.delete()
        revision.generated_document.refresh_from_db()
        self.assertEqual(revision.generated_document.status, "current")

        self.document.file_name = "changed-selected-demand.docx"
        self.document.save()
        revision.generated_document.refresh_from_db()
        self.assertEqual(revision.generated_document.status, "out_of_date")

        approve_document_revision(revision, self.user)
        self.document.delete()
        revision.generated_document.refresh_from_db()
        self.assertEqual(revision.generated_document.status, "out_of_date")

    @patch("project_content.services.requests.post")
    def test_fat_cat_are_independent_and_manual_edits_survive_hints(self, post):
        post.return_value = SimpleNamespace(status_code=200, json=self.acceptance_result)
        run, _ = create_acceptance_generation_run(
            project=self.project,
            user=self.user,
            document_ids=[self.document.id],
        )
        process_acceptance_generation_run(run.id)
        run.refresh_from_db()
        apply_acceptance_generation_run(run, self.user)

        criteria = AcceptanceCriterion.objects.filter(checklist__project=self.project)
        self.assertEqual(criteria.count(), 2)
        self.assertEqual(
            set(criteria.values_list("checklist__phase", flat=True)),
            {AcceptanceChecklist.Phase.FACTORY, AcceptanceChecklist.Phase.SITE},
        )
        self.assertFalse(
            criteria.filter(evidence__generation_run__isnull=True).exists()
        )
        fat = criteria.get(checklist__phase=AcceptanceChecklist.Phase.FACTORY)
        fat.acceptance_criteria = "Manually approved test wording."
        fat.manual_fields = ["acceptance_criteria"]
        fat.status = "Passed"
        fat.save()

        self.requirement.custom_fields = {"value": "400 V", "inspection": "witnessed"}
        self.requirement.save()
        post.return_value = SimpleNamespace(
            status_code=200,
            json=lambda: self.acceptance_result("modify"),
        )
        hint_run, _ = create_acceptance_generation_run(
            project=self.project,
            user=self.user,
            document_ids=[self.document.id],
        )
        self.assertTrue(hint_run.hint_only)
        process_acceptance_generation_run(hint_run.id)
        hint_run.refresh_from_db()
        apply_acceptance_generation_run(
            hint_run,
            self.user,
            selected_keys=["safety-test"],
        )
        fat.refresh_from_db()
        self.assertEqual(fat.acceptance_criteria, "Manually approved test wording.")
        self.assertEqual(fat.status, "Passed")

    @override_settings(PROJECT_CONTENT_AI_TIMEOUT_SECONDS=137)
    @patch("project_content.services.requests.post")
    def test_acceptance_generation_omits_unconfigured_ai_priority(self, post):
        configuration, _ = ECPMPConfig.objects.get_or_create(project=self.project)
        configuration.priority_levels = ["High:#fee2e2", "Low:#dbeafe"]
        configuration.save(update_fields=["priority_levels"])
        result = self.acceptance_result()
        result["criteria"][0]["priority"] = "Critical"
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)

        run, _ = create_acceptance_generation_run(
            project=self.project,
            user=self.user,
            document_ids=[self.document.id],
        )
        process_acceptance_generation_run(run.id)
        run.refresh_from_db()

        self.assertEqual(run.status, AcceptanceGenerationRun.Status.REVIEW_READY)
        self.assertIsNone(run.draft_content["criteria"][0]["priority"])
        self.assertIn('Priority "Critical" was omitted', run.warnings[0])
        self.assertEqual(
            post.call_args.kwargs["json"]["options"]["allowed_priorities"],
            ["High", "Low"],
        )
        self.assertEqual(post.call_args.kwargs["timeout"], 137)

    def test_source_change_rejects_review_apply(self):
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        run.status = ProjectDocumentGenerationRun.Status.REVIEW_READY
        run.draft_content = self.offer_result()["content"]
        run.save()
        self.requirement.name = "Updated voltage requirement"
        self.requirement.save()
        with self.assertRaises(ProjectContentError) as context:
            apply_document_generation_run(run, self.user)
        self.assertEqual(context.exception.code, "source_changed")

    @patch("project_content.services.requests.post")
    @patch("project_content.rendering.convert_docx_to_pdf", return_value=b"%PDF-1.7\n")
    def test_acceptance_apply_rejects_changed_approved_document_input(self, _pdf, post):
        self.approved_specification_fixture()
        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        document_run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        process_document_generation_run(document_run.id)
        document_run.refresh_from_db()
        revision = apply_document_generation_run(document_run, self.user)
        approve_document_revision(revision, self.user)

        acceptance_run, _ = create_acceptance_generation_run(
            project=self.project,
            user=self.user,
            document_ids=[self.document.id],
        )
        acceptance_run.status = AcceptanceGenerationRun.Status.REVIEW_READY
        acceptance_run.draft_content = self.acceptance_result()
        acceptance_run.save(update_fields=["status", "draft_content", "updated_at"])

        revision.approval_status = revision.ApprovalStatus.SUPERSEDED
        revision.save(update_fields=["approval_status"])
        with self.assertRaises(ProjectContentError) as context:
            apply_acceptance_generation_run(acceptance_run, self.user)
        self.assertEqual(context.exception.code, "source_changed")

    def test_generation_snapshot_supports_selected_process_and_opportunity_sources(self):
        opportunity = Opportunity.objects.create(
            created_by=self.user,
            company_profile=self.company,
            title="Verified commercial source",
            details="Approved proposal context.",
            value=125000,
            planned_delivery_start_date="2026-10-01",
        )
        opportunity.projects.add(self.project)

        run, created = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[],
            request_options={
                "structured_sources": ["process_stations", "opportunities"]
            },
        )

        self.assertTrue(created)
        self.assertEqual(len(run.source_snapshot["process_stations"]), 1)
        self.assertEqual(run.source_snapshot["opportunities"][0]["value"], "125000.0")
        self.assertNotIn("details", run.source_snapshot["opportunities"][0])
        self.assertNotIn("probability", run.source_snapshot["opportunities"][0])
        self.assertNotIn("sales_phase", run.source_snapshot["opportunities"][0])
        self.assertNotIn("status", run.source_snapshot["opportunities"][0])
        self.assertEqual(run.source_snapshot["customer_demands"], [])
        self.assertEqual(run.source_snapshot["bill_of_material"], [])
        self.assertEqual(
            run.source_snapshot["generation_contract"],
            {"kind": "offer", "version": "4"},
        )

    def test_generation_contract_change_does_not_reuse_legacy_offer_draft(self):
        first_run, created = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        self.assertTrue(created)
        first_run.status = ProjectDocumentGenerationRun.Status.REVIEW_READY
        legacy_snapshot = dict(first_run.source_snapshot)
        legacy_snapshot["generation_contract"] = {"kind": "offer", "version": "3"}
        first_run.source_snapshot = legacy_snapshot
        first_run.source_hash = canonical_hash(legacy_snapshot)
        first_run.save(
            update_fields=["status", "source_snapshot", "source_hash", "updated_at"]
        )

        current_run, created = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )

        self.assertTrue(created)
        self.assertNotEqual(current_run.pk, first_run.pk)
        self.assertEqual(
            current_run.source_snapshot["generation_contract"],
            {"kind": "offer", "version": "4"},
        )

    def test_generation_rejects_unknown_structured_source_selection(self):
        with self.assertRaises(ProjectContentError) as context:
            create_document_generation_run(
                project=self.project,
                user=self.user,
                kind="offer",
                request_options={"structured_sources": ["private_database"]},
            )
        self.assertEqual(context.exception.code, "invalid_structured_sources")

    def test_grounding_accepts_only_explicitly_labeled_source_ids(self):
        opportunity = Opportunity.objects.create(
            created_by=self.user,
            company_profile=self.company,
            title="Verified commercial source",
            value=125000,
        )
        opportunity.projects.add(self.project)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[],
            request_options={"structured_sources": ["opportunities"]},
        )
        evidence = [
            {
                "source_type": "opportunity",
                "source_id": str(opportunity.id),
                "quote": "125000.0",
            }
        ]

        _validate_result_references(
            run.source_snapshot,
            {
                "content": {
                    "executive_summary": [
                        {
                            "text": f"Opportunity {opportunity.id} lists value 125000.0.",
                            "evidence": evidence,
                        }
                    ]
                }
            },
        )
        with self.assertRaises(ProjectContentError) as context:
            _validate_result_references(
                run.source_snapshot,
                {
                    "content": {
                        "executive_summary": [
                            {
                                "text": f"The supply voltage is {opportunity.id} V.",
                                "evidence": evidence,
                            }
                        ]
                    }
                },
            )
        self.assertEqual(context.exception.code, "unsupported_quantitative_value")

    def test_grounding_normalizes_unique_requirement_business_ids(self):
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="requirements_specification",
            document_ids=[self.document.id],
            request_options={"structured_sources": ["requirements"]},
        )
        result = {
            "content": {
                "traceability_matrix": [
                    {
                        "requirement_id": "REQ-001",
                        "coverage": "The supply requirement is 400 V.",
                        "evidence": [
                            {
                                "source_type": "requirement",
                                "source_id": "REQ-001",
                                "requirement_id": "REQ-001",
                                "quote": "400 V",
                            }
                        ],
                    }
                ],
                "omitted_requirement_ids": [],
            }
        }

        _normalize_result_references(run.source_snapshot, result)

        coverage = result["content"]["traceability_matrix"][0]
        self.assertEqual(coverage["requirement_id"], str(self.requirement.id))
        self.assertEqual(
            coverage["evidence"][0]["requirement_id"],
            str(self.requirement.id),
        )
        self.assertEqual(
            coverage["evidence"][0]["source_id"],
            str(self.requirement.id),
        )
        _validate_result_references(run.source_snapshot, result)

    def test_grounding_discards_requirement_ids_from_document_evidence(self):
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="requirements_specification",
            document_ids=[self.document.id],
            request_options={"structured_sources": ["requirements"]},
        )
        result = {
            "content": {
                "traceability_matrix": [
                    {
                        "requirement_id": str(self.requirement.id),
                        "coverage": "The supply requirement is 400 V.",
                        "evidence": [
                            {
                                "document_id": str(self.document.id),
                                "segment_id": str(self.segment.id),
                                "source_type": "document",
                                "source_id": "CDE-RS-001",
                                "requirement_id": "CDE-RS-001",
                                "quote": "400 V",
                            }
                        ],
                    }
                ],
                "omitted_requirement_ids": [],
            }
        }

        _normalize_result_references(run.source_snapshot, result)

        evidence = result["content"]["traceability_matrix"][0]["evidence"][0]
        self.assertIsNone(evidence["source_id"])
        self.assertIsNone(evidence["requirement_id"])
        _validate_result_references(run.source_snapshot, result)
        _validate_requirements_specification_coverage(run.source_snapshot, result)

    def test_grounding_repairs_only_unique_exact_document_quote_matches(self):
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        result = self.offer_result()
        evidence = result["content"]["executive_summary"][0]["evidence"][0]
        evidence["segment_id"] = "missing"

        _normalize_result_references(run.source_snapshot, result)

        self.assertEqual(evidence["document_id"], str(self.document.id))
        self.assertEqual(evidence["segment_id"], str(self.segment.id))
        _validate_result_references(run.source_snapshot, result)

    @patch("project_content.services.requests.post")
    def test_offer_rejects_commercial_values_from_document_evidence(self, post):
        result = self.offer_result()
        result["content"]["commercial_values"] = [
            {
                "text": "The commercial value is EUR 400.",
                "evidence": [
                    {
                        "document_id": str(self.document.id),
                        "segment_id": str(self.segment.id),
                        "source_type": "document",
                        "source_id": None,
                        "requirement_id": None,
                        "location": "Page 1",
                        "quote": "400 V",
                    }
                ],
                "unresolved": False,
            }
        ]
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )

        process_document_generation_run(run.id)
        run.refresh_from_db()

        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.FAILED)
        self.assertEqual(run.error_code, "unsupported_commercial_source")

    @patch("project_content.services.requests.post")
    def test_offer_rejects_internal_opportunity_metadata(self, post):
        opportunity = Opportunity.objects.create(
            created_by=self.user,
            company_profile=self.company,
            title="Internal sales context",
            value=125000,
            probability=80,
        )
        opportunity.projects.add(self.project)
        result = self.offer_result()
        result["content"]["commercial_values"] = [
            {
                "text": "The sales probability is 80 percent.",
                "evidence": [
                    {
                        "document_id": None,
                        "segment_id": None,
                        "source_type": "opportunity",
                        "source_id": str(opportunity.id),
                        "requirement_id": None,
                        "location": "probability",
                        "quote": "80.0",
                    }
                ],
                "unresolved": False,
            }
        ]
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
            request_options={"structured_sources": ["opportunities"]},
        )

        process_document_generation_run(run.id)
        run.refresh_from_db()

        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.FAILED)
        self.assertEqual(run.error_code, "internal_opportunity_metadata")

    @patch("project_content.services.requests.post")
    def test_offer_requires_evidence_for_requirement_coverage(self, post):
        result = self.offer_result()
        result["content"]["requirements_coverage"] = [
            {
                "requirement_id": str(self.requirement.id),
                "coverage": "Supply voltage is covered.",
                "evidence": [],
            }
        ]
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )

        process_document_generation_run(run.id)
        run.refresh_from_db()

        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.FAILED)
        self.assertEqual(run.error_code, "unsupported_offer_coverage")

    @patch("project_content.services.requests.post")
    def test_offer_requires_explicit_source_conflicts_to_remain_visible(self, post):
        DocumentContentSegment.objects.create(
            document=self.document,
            index=self.index,
            segment_type=DocumentContentSegment.SegmentType.PAGE,
            sequence=2,
            location_label="Revision notes",
            text=(
                "A 60 second cycle is not approved. Do not choose between the "
                "45 and 60 second values; report the conflict for review."
            ),
        )
        result = self.offer_result()
        result["content"].pop("open_questions")
        result["warnings"] = []
        post.return_value = SimpleNamespace(status_code=200, json=lambda: result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )

        process_document_generation_run(run.id)
        run.refresh_from_db()

        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.FAILED)
        self.assertEqual(run.error_code, "unresolved_source_conflict")

    def test_default_template_contains_controlled_marker(self):
        docx_bytes = build_docx_bytes(
            "offer",
            {"sections": [{"title": "Scope", "content": ["Supply machine"]}]},
        )
        self.assertGreater(len(docx_bytes), 500)
        self.assertNotIn(CONTENT_MARKER.encode(), docx_bytes)

    def test_renderer_version_prevents_reusing_legacy_artifacts(self):
        template = SimpleNamespace(id=17, version=4)

        self.assertEqual(
            document_template_version(template),
            "company-17-v4-renderer-v4",
        )
        self.assertEqual(
            document_template_version(None),
            "system-v1-renderer-v4",
        )

    def test_offer_renderer_uses_business_order_and_human_readable_sources(self):
        from docx import Document as DocxDocument

        content = {
            "exclusions": [
                {
                    "text": "Civil works are excluded.",
                    "evidence": [],
                    "unresolved": True,
                }
            ],
            "executive_summary": [
                {
                    "text": "Supply the DM-200 drilling machine.",
                    "evidence": [
                        {
                            "document_id": str(self.document.id),
                            "segment_id": str(self.segment.id),
                            "source_type": "document",
                            "location": "Page 1",
                            "quote": "REQ-001 requires a 400 V supply",
                        }
                    ],
                    "unresolved": False,
                }
            ],
            "requirements_coverage": [
                {
                    "requirement_id": str(self.requirement.id),
                    "coverage": "The 400 V supply requirement is covered.",
                    "evidence": [
                        {
                            "source_type": "requirement",
                            "source_id": str(self.requirement.id),
                            "requirement_id": str(self.requirement.id),
                            "location": "value",
                            "quote": "400 V",
                        }
                    ],
                }
            ],
            "open_questions": ["Confirm the commercial price."],
        }
        snapshot, _ = build_project_content_snapshot(self.project)

        docx_bytes = build_docx_bytes(
            "offer",
            content,
            locale="en",
            source_snapshot=snapshot,
            revision=2,
        )
        document = DocxDocument(BytesIO(docx_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        complete_text = "\n".join(
            paragraphs
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )

        self.assertLess(
            paragraphs.index("Executive Summary"),
            paragraphs.index("Exclusions"),
        )
        self.assertIn("Document Information", paragraphs)
        self.assertIn("Source References", paragraphs)
        self.assertIn("DM-200 Demand.docx", complete_text)
        self.assertIn("REQ-001 - Supply voltage", complete_text)
        self.assertIn("Unresolved: Civil works are excluded.", complete_text)
        self.assertNotIn("Segment Id", complete_text)
        self.assertNotIn("Document Id", complete_text)
        self.assertNotIn("Source Type", complete_text)
        self.assertNotIn("Unresolved\nTrue", complete_text)

    def test_requirements_specification_renderer_uses_canonical_sections(self):
        from docx import Document as DocxDocument

        snapshot, _ = build_project_content_snapshot(self.project)
        content = {
            "functional_requirements": [
                {
                    "text": "The machine shall operate at 400 V.",
                    "evidence": [
                        {
                            "document_id": str(self.document.id),
                            "segment_id": str(self.segment.id),
                            "source_type": "document",
                            "location": "Page 1",
                            "quote": "REQ-001 requires a 400 V supply",
                        }
                    ],
                    "unresolved": False,
                }
            ],
            "purpose_and_scope": [
                {
                    "text": "Specify the DM-200 supply requirements.",
                    "evidence": [
                        {
                            "document_id": str(self.document.id),
                            "segment_id": str(self.segment.id),
                            "source_type": "document",
                            "location": "Page 1",
                            "quote": "REQ-001 requires a 400 V supply",
                        }
                    ],
                    "unresolved": False,
                }
            ],
            "traceability_matrix": [
                {
                    "requirement_id": str(self.requirement.id),
                    "coverage": "Covered by the functional supply requirement.",
                    "evidence": [
                        {
                            "source_type": "requirement",
                            "source_id": str(self.requirement.id),
                            "requirement_id": str(self.requirement.id),
                            "location": "value",
                            "quote": "400 V",
                        }
                    ],
                }
            ],
            "unresolved_questions": ["Confirm the applicable safety standard."],
        }

        docx_bytes = build_docx_bytes(
            "requirements_specification",
            content,
            locale="en",
            source_snapshot=snapshot,
            revision=1,
        )
        document = DocxDocument(BytesIO(docx_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        complete_text = "\n".join(
            paragraphs
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )

        self.assertLess(
            paragraphs.index("Purpose and Scope"),
            paragraphs.index("Functional Requirements"),
        )
        self.assertIn("Requirements Traceability", paragraphs)
        self.assertIn("REQ-001 - Supply voltage", complete_text)
        self.assertIn("DM-200 Demand.docx", complete_text)
        self.assertNotIn("Segment Id", complete_text)
        self.assertNotIn("Document Id", complete_text)

    def test_specification_document_preserves_four_workstation_levels(self):
        from docx import Document as DocxDocument

        content = {
            "workstation_hierarchy": [
                {"text": name, "level": level, "evidence": []}
                for level, name in enumerate(
                    ["Station", "Assembly", "Subassembly", "Purchased part"], 1
                )
            ],
        }
        document = DocxDocument(BytesIO(build_docx_bytes(
            "requirements_specification", content, locale="en"
        )))
        paragraphs = {
            paragraph.text: paragraph for paragraph in document.paragraphs
        }
        indents = [
            paragraphs[name].paragraph_format.left_indent
            for name in ["Station", "Assembly", "Subassembly", "Purchased part"]
        ]
        self.assertTrue(all(indent is not None for indent in indents), indents)
        self.assertTrue(all(a < b for a, b in zip(indents, indents[1:])), indents)

    def test_specification_orders_shuffled_occurrences_parent_first_without_merging_stations(self):
        from docx import Document as DocxDocument
        from .services import _normalize_result_references

        rows = [
            ("a", "leaf", "child", "Part A", 3, 0),
            ("b", "child", "root", "Part B", 2, 0),
            ("a", "child", "root", "Assembly A", 2, 0),
            ("a", "root", None, "Station A", 1, 0),
            ("b", "root", None, "Station B", 1, 0),
        ]
        content = {"workstation_hierarchy": [
            dict(workstation_id=station, occurrence_id=node, parent_occurrence_id=parent,
                 text=name, level=level, sort_order=order, evidence=[])
            for station, node, parent, name, level, order in rows
        ]}
        expected = ["Station A", "Assembly A", "Part A", "Station B", "Part B"]
        document = DocxDocument(BytesIO(build_docx_bytes(
            "requirements_specification", content, locale="en"
        )))
        self.assertEqual([p.text for p in document.paragraphs if p.text in expected], expected)
        _normalize_result_references({}, {"content": content})
        self.assertEqual([row["text"] for row in content["workstation_hierarchy"]], expected)

    def test_document_reference_headers_do_not_split_and_have_readable_widths(self):
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn
        from docx.shared import Inches

        document = DocxDocument(BytesIO(build_docx_bytes(
            "requirements_specification",
            {"traceability_matrix": [{
                "requirement_id": str(self.requirement.id),
                "coverage": "Covered by supplied evidence.",
                "evidence": [{"source_type": "requirement", "source_id": str(self.requirement.id), "quote": "400 V"}],
            }]},
            locale="en",
        )))
        traceability, references = document.tables
        for table in document.tables:
            properties = table.rows[0]._tr.get_or_add_trPr()
            self.assertIsNotNone(properties.find(qn("w:tblHeader")))
            self.assertIsNotNone(properties.find(qn("w:cantSplit")))
        self.assertGreaterEqual(traceability.columns[2].width, Inches(1.1))
        self.assertGreaterEqual(references.columns[0].width, Inches(1.2))

    def test_offer_attachment_table_includes_localized_type_and_readable_revision(self):
        from docx import Document as DocxDocument
        from docx.shared import Inches

        for locale, expected_type, expected_status in [
            ("en", "Technical Specification", "Draft"),
            ("de", "Technische Spezifikation", "Entwurf"),
            ("fr", "Spécification technique", "Brouillon"),
        ]:
            with self.subTest(locale=locale):
                document = DocxDocument(BytesIO(build_docx_bytes(
                    "offer", {"attachment_manifest": [{
                        "name": "Battery-Technical-Specification-R3.pdf",
                        "document_type": "technical_specification",
                        "version": 3,
                        "generation_status": "draft",
                    }]}, locale=locale,
                )))
                table = document.tables[0]
                self.assertEqual(len(table.columns), 4)
                self.assertEqual(
                    [cell.text for cell in table.rows[1].cells],
                    ["Battery-Technical-Specification-R3.pdf", expected_type, "3", expected_status],
                )
                self.assertGreaterEqual(table.columns[2].width, Inches(1.0))
                self.assertGreaterEqual(table.columns[3].width, Inches(1.2))
                self.assertLessEqual(sum(column.width for column in table.columns), Inches(6.5))

    def test_downloadable_system_template_contains_controlled_marker(self):
        import io
        import zipfile

        from .rendering import system_template_bytes

        with zipfile.ZipFile(io.BytesIO(system_template_bytes("offer"))) as archive:
            self.assertIn(CONTENT_MARKER.encode(), archive.read("word/document.xml"))

    def test_offer_attachment_approval_states_are_localized(self):
        from docx import Document as DocxDocument

        for locale, expected in [
            ("en", ["Approved", "Superseded"]),
            ("de", ["Freigegeben", "Ersetzt"]),
            ("fr", ["Approuvé", "Remplacé"]),
        ]:
            with self.subTest(locale=locale):
                document = DocxDocument(BytesIO(build_docx_bytes(
                    "offer", {"attachment_manifest": [
                        {"name": "Specification.pdf", "document_type": "technical_specification",
                         "version": 3, "generation_status": state}
                        for state in ("approved", "superseded")
                    ]}, locale=locale,
                )))
                self.assertEqual(
                    [row.cells[3].text for row in document.tables[0].rows[1:]], expected,
                )

    def test_template_validator_rejects_non_docx(self):
        template = SimpleNamespace(
            open=lambda mode: None,
            read=lambda: b"not a docx",
            close=lambda: None,
        )
        with self.assertRaises(TemplateValidationError) as context:
            validate_template_file(template, "template.docm")
        self.assertEqual(context.exception.code, "unsupported_template_type")

    def test_template_validator_rejects_missing_marker_and_external_relationship(self):
        from .rendering import system_template_bytes

        source = system_template_bytes("offer")
        without_marker = BytesIO()
        with zipfile.ZipFile(BytesIO(source)) as source_archive, zipfile.ZipFile(
            without_marker, "w"
        ) as output_archive:
            for entry in source_archive.infolist():
                value = source_archive.read(entry)
                if entry.filename == "word/document.xml":
                    value = value.replace(CONTENT_MARKER.encode(), b"CONTENT")
                output_archive.writestr(entry, value)
        with self.assertRaises(TemplateValidationError) as context:
            validate_template_file(
                ContentFile(without_marker.getvalue(), name="missing-marker.docx"),
                "missing-marker.docx",
            )
        self.assertEqual(context.exception.code, "missing_content_marker")

        external_relationship = BytesIO()
        with zipfile.ZipFile(BytesIO(source)) as source_archive, zipfile.ZipFile(
            external_relationship, "w"
        ) as output_archive:
            for entry in source_archive.infolist():
                value = source_archive.read(entry)
                if entry.filename == "word/_rels/document.xml.rels":
                    value = value.replace(
                        b"</Relationships>",
                        (
                            b'<Relationship Id="external" Type="reference" '
                            b'Target="https://example.com/template" '
                            b'TargetMode="External"/></Relationships>'
                        ),
                    )
                output_archive.writestr(entry, value)
        with self.assertRaises(TemplateValidationError) as context:
            validate_template_file(
                ContentFile(
                    external_relationship.getvalue(),
                    name="external-relationship.docx",
                ),
                "external-relationship.docx",
            )
        self.assertEqual(context.exception.code, "external_template_relationship")

    def test_template_validator_rejects_archive_traversal(self):
        from .rendering import system_template_bytes

        unsafe = BytesIO()
        with zipfile.ZipFile(BytesIO(system_template_bytes("offer"))) as source_archive, zipfile.ZipFile(
            unsafe, "w"
        ) as output_archive:
            for entry in source_archive.infolist():
                output_archive.writestr(entry, source_archive.read(entry))
            output_archive.writestr("../outside.txt", "unsafe")
        with self.assertRaises(TemplateValidationError) as context:
            validate_template_file(
                ContentFile(unsafe.getvalue(), name="unsafe.docx"),
                "unsafe.docx",
            )
        self.assertEqual(context.exception.code, "unsafe_template_path")

    def test_system_folder_cannot_be_renamed_through_serializer(self):
        ensure_project_content_folders(self.project)
        folder = Folder.objects.get(project=self.project, system_key="proposal")
        serializer = FolderSerializer(
            folder,
            data={"name": "Repurposed"},
            partial=True,
            context={"request": SimpleNamespace(user=self.user), "user": self.user},
        )
        self.assertFalse(serializer.is_valid())
        self.assertIn("cannot be renamed", str(serializer.errors))

    def test_acceptance_assignees_must_belong_to_project(self):
        outsider = User.objects.create_user(
            email="acceptance-outsider@example.com",
            password="password",
            first_name="Acceptance",
            last_name="Outsider",
            user_type="company_user",
            company=self.company,
        )
        checklist = AcceptanceChecklist.objects.create(
            project=self.project,
            phase=AcceptanceChecklist.Phase.FACTORY,
            display_name="Factory Acceptance Test (FAT)",
            created_by=self.user,
        )
        serializer = AcceptanceCriterionSerializer(
            data={
                "checklist": str(checklist.id),
                "name": "Unauthorized assignment",
                "acceptance_criteria": "The project team approves the result.",
                "assignees": [outsider.id],
            },
            context={"request": SimpleNamespace(user=self.user)},
        )
        self.assertFalse(serializer.is_valid())
        self.assertIn("Assignees must belong", str(serializer.errors))

    @patch("ecpmpapp.permissions.LicensePermission.has_permission", return_value=True)
    @patch("project_content.views.require_project_content_access")
    def test_acceptance_list_filters_paginate_and_return_hierarchy_options(
        self,
        _require_project_content_access,
        _license_permission,
    ):
        configuration, _ = ECPMPConfig.objects.get_or_create(project=self.project)
        configuration.acceptance_status_options = ["Passed:#d1fae5"]
        configuration.priority_levels = ["High:#fee2e2"]
        configuration.labels = ["Safety:#fef3c7"]
        configuration.save()
        checklist = AcceptanceChecklist.objects.create(
            project=self.project,
            phase=AcceptanceChecklist.Phase.FACTORY,
            display_name="Factory Acceptance Test (FAT)",
            created_by=self.user,
        )
        criterion = AcceptanceCriterion.objects.create(
            checklist=checklist,
            stable_key="filtered-safety-test",
            name="Filtered safety test",
            acceptance_criteria="The safety circuit passes.",
            start_date=date(2026, 9, 1),
            status="Passed",
            priority="High",
            labels=["Safety"],
            created_by=self.user,
            modified_by=self.user,
        )
        criterion.assignees.add(self.user)

        client = APIClient()
        client.force_authenticate(self.user)
        response = client.get(
            "/api/project-content/acceptance-criteria/",
            {
                "project": self.project.id,
                "checklist": str(checklist.id),
                "status": "Passed",
                "priority": "High",
                "assignee": self.user.id,
                "label": "Safety",
                "date_field": "start_date",
                "date_from": "2026-09-01",
                "date_to": "2026-09-30",
                "page": 1,
                "page_size": 10,
            },
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["results"][0]["id"], str(criterion.id))

        response = client.get(
            "/api/project-content/acceptance-criteria/options/",
            {"project": self.project.id, "checklist": str(checklist.id)},
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.data[0]["id"], str(criterion.id))

        response = client.post(
            "/api/project-content/acceptance-criteria/",
            {
                "checklist": str(checklist.id),
                "name": "Invalid dates",
                "acceptance_criteria": "Date validation",
                "start_date": "2026-09-10",
                "end_date": "2026-09-01",
                "status": "Passed",
                "priority": "High",
                "labels": ["Safety"],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("end_date", response.data)

    def test_acceptance_comment_rejects_cross_project_attachments(self):
        checklist = AcceptanceChecklist.objects.create(
            project=self.project,
            phase=AcceptanceChecklist.Phase.FACTORY,
            display_name="Factory Acceptance Test (FAT)",
            created_by=self.user,
        )
        criterion = AcceptanceCriterion.objects.create(
            checklist=checklist,
            stable_key="manual-comment-test",
            name="Comment test",
            acceptance_criteria="The evidence is reviewed.",
            created_by=self.user,
            modified_by=self.user,
        )
        other_project = Project.objects.create(
            project_name="DM-201",
            project_description="Other project",
            contract_value=1000,
            company_profile=self.company,
            user=self.user,
            created_by=self.user,
        )
        other_root = Folder.objects.create(
            project=other_project,
            name=f"{other_project.id}_root_folder",
        )
        other_document = Document.objects.create(
            folder=other_root,
            author=self.user,
            created_by=self.user,
            file_name="Other Project Evidence.pdf",
            status="Draft",
        )
        serializer = AcceptanceCriterionCommentSerializer(
            data={
                "criterion": str(criterion.id),
                "content": "This attachment must not cross project boundaries.",
                "attachments": [other_document.id],
            },
            context={"request": SimpleNamespace(user=self.user)},
        )
        self.assertFalse(serializer.is_valid())
        self.assertIn("attachments must belong", str(serializer.errors))

    def test_demo_seed_and_verification_commands_are_repeatable(self):
        output = StringIO()
        call_command(
            "seed_project_content_demo",
            project_id=self.project.id,
            apply=True,
            include_edge_documents=True,
            stdout=output,
        )
        call_command(
            "seed_project_content_demo",
            project_id=self.project.id,
            apply=True,
            include_edge_documents=True,
            stdout=output,
        )
        call_command(
            "verify_project_content_demo",
            project_id=self.project.id,
            minimum_criteria=12,
            require_edge_documents=True,
            stdout=output,
        )
        self.assertIn("demo verified", output.getvalue())

    def test_demo_verification_accepts_parser_segment_recomposition(self):
        output = StringIO()
        call_command(
            "seed_project_content_demo",
            project_id=self.project.id,
            apply=True,
            include_edge_documents=True,
            stdout=output,
        )
        edge_documents = Document.objects.filter(
            folder__project=self.project,
            folder__name="Project Content AI Edge Sources",
        ).select_related("search_index")
        for document in edge_documents:
            text = "\n".join(
                document.content_segments.order_by("sequence", "id").values_list(
                    "text",
                    flat=True,
                )
            )
            document.content_segments.all().delete()
            DocumentContentSegment.objects.create(
                document=document,
                index=document.search_index,
                segment_type=DocumentContentSegment.SegmentType.DOCUMENT,
                sequence=1,
                location_label="Document",
                text=text,
            )

        call_command(
            "verify_project_content_demo",
            project_id=self.project.id,
            minimum_criteria=12,
            require_edge_documents=True,
            stdout=output,
        )
        self.assertIn("demo verified", output.getvalue())

    def test_performance_seed_dry_run_reports_full_fixture_size(self):
        output = StringIO()
        call_command(
            "seed_project_content_demo",
            project_id=self.project.id,
            include_performance=True,
            stdout=output,
        )
        self.assertIn("500 requirements", output.getvalue())
        self.assertIn("2000 FAT/CAT criteria", output.getvalue())

    def test_edge_document_seed_dry_run_reports_mixed_quality_pack(self):
        output = StringIO()
        call_command(
            "seed_project_content_demo",
            project_id=self.project.id,
            include_edge_documents=True,
            stdout=output,
        )
        self.assertIn("16 source documents", output.getvalue())

    def test_tree_snapshot_numbers_hierarchy_in_one_query(self):
        child = Requirement.objects.create(
            project=self.project,
            parent=self.requirement,
            name="Child requirement",
            object_number="REQ-001-1",
            sort_order=1,
            level=2,
            created_by=self.user,
            modified_by=self.user,
        )
        with self.assertNumQueries(1):
            rows = _tree_rows(
                Requirement.objects.filter(project=self.project),
                "name",
                "object_number",
            )
        hierarchy_numbers = {row["id"]: row["hierarchy_number"] for row in rows}
        self.assertEqual(hierarchy_numbers[str(self.requirement.id)], "1")
        self.assertEqual(hierarchy_numbers[str(child.id)], "1.1")

    @patch("project_content.services.requests.post")
    def test_generation_lifecycle_creates_redacted_audit_events(self, post):
        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        process_document_generation_run(run.id)
        event_types = set(
            ProjectContentAuditEvent.objects.filter(project=self.project).values_list(
                "event_type", flat=True
            )
        )
        self.assertIn("document_generation_requested", event_types)
        self.assertIn("document_generation_review_ready", event_types)
        self.assertFalse(
            ProjectContentAuditEvent.objects.filter(
                metadata__has_key="source_snapshot"
            ).exists()
        )

    @patch("project_content.services.requests.post")
    def test_retryable_document_failure_returns_to_queued_state(self, post):
        post.return_value = SimpleNamespace(status_code=502)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )

        run = process_document_generation_run(run.id)
        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.FAILED)

        _mark_generation_retry_pending(run)
        run.refresh_from_db()
        run.generated_document.refresh_from_db()
        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.QUEUED)
        self.assertEqual(run.error_code, "")
        self.assertIsNone(run.completed_at)
        self.assertEqual(
            run.generated_document.status,
            GeneratedProjectDocument.Status.GENERATING,
        )

    @patch("project_content.services.requests.post")
    def test_retryable_acceptance_failure_returns_to_queued_state(self, post):
        post.return_value = SimpleNamespace(status_code=502)
        run, _ = create_acceptance_generation_run(
            project=self.project,
            user=self.user,
            document_ids=[self.document.id],
        )

        run = process_acceptance_generation_run(run.id)
        self.assertEqual(run.status, AcceptanceGenerationRun.Status.FAILED)

        _mark_generation_retry_pending(run)
        run.refresh_from_db()
        self.assertEqual(run.status, AcceptanceGenerationRun.Status.QUEUED)
        self.assertEqual(run.error_code, "")
        self.assertIsNone(run.completed_at)

    @patch("project_content.services.dispatch_document_render")
    @patch("project_content.services.requests.post")
    def test_failed_document_render_can_be_requeued_without_ai_regeneration(
        self,
        post,
        dispatch_document_render,
    ):
        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        run, _ = create_document_generation_run(
            project=self.project,
            user=self.user,
            kind="offer",
            document_ids=[self.document.id],
        )
        run = process_document_generation_run(run.id)
        run.status = ProjectDocumentGenerationRun.Status.FAILED
        run.error_code = "document_render_failed"
        run.error_message = "The reviewed document could not be rendered."
        run.save(update_fields=["status", "error_code", "error_message", "updated_at"])

        with self.captureOnCommitCallbacks(execute=True):
            run = queue_document_generation_run(run, self.user)

        run.refresh_from_db()
        run.generated_document.refresh_from_db()
        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.APPLYING)
        self.assertEqual(run.error_code, "")
        self.assertEqual(
            run.generated_document.status,
            GeneratedProjectDocument.Status.GENERATING,
        )
        dispatch_document_render.assert_called_once_with(run.id, self.user.id)
        self.assertTrue(
            ProjectContentAuditEvent.objects.filter(
                project=self.project,
                event_type="document_render_retry_requested",
                object_id=str(run.id),
            ).exists()
        )

    @override_settings(PROJECT_CONTENT_RUN_SYNCHRONOUSLY=True)
    @patch("ecpmpapp.permissions.LicensePermission.has_permission", return_value=True)
    @patch("project_content.views.require_project_content_access")
    @patch("project_content.rendering.convert_docx_to_pdf", return_value=b"%PDF-1.7\n")
    @patch("project_content.services.requests.post")
    def test_project_content_http_workflow(
        self,
        post,
        _pdf,
        _require_project_content_access,
        _license_permission,
    ):
        self.approved_specification_fixture()
        client = APIClient()
        client.force_authenticate(self.user)

        post.return_value = SimpleNamespace(status_code=200, json=self.offer_result)
        response = client.post(
            f"/api/project-content/projects/{self.project.id}/offer/generate/",
            {"locale": "en"},
            format="json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        self.assertEqual(response.data["status"], "review_ready")
        run_id = response.data["id"]

        reviewed_content = {
            **response.data["draft_content"],
            "open_questions": ["Confirm delivery date."],
        }
        response = client.patch(
            f"/api/project-content/generation-runs/{run_id}/",
            {"content": reviewed_content},
            format="json",
        )
        self.assertEqual(response.status_code, 200)

        with self.captureOnCommitCallbacks(execute=True):
            response = client.post(
                f"/api/project-content/generation-runs/{run_id}/apply/",
                {},
                format="json",
            )
        self.assertEqual(response.status_code, 202, response.content)
        run = ProjectDocumentGenerationRun.objects.get(pk=run_id)
        self.assertEqual(run.status, ProjectDocumentGenerationRun.Status.APPLIED)
        revision = run.applied_revision

        response = client.post(
            f"/api/project-content/revisions/{revision.id}/approve/",
            {},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["approval_status"], "approved")

        post.return_value = SimpleNamespace(status_code=200, json=self.acceptance_result)
        response = client.post(
            f"/api/project-content/projects/{self.project.id}/acceptance/generate/",
            {"locale": "en"},
            format="json",
        )
        self.assertEqual(response.status_code, 202, response.content)
        self.assertEqual(response.data["status"], "review_ready")
        acceptance_run_id = response.data["id"]

        response = client.post(
            f"/api/project-content/acceptance-runs/{acceptance_run_id}/apply/",
            {},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["status"], "applied")

        fat = AcceptanceCriterion.objects.get(
            checklist__project=self.project,
            checklist__phase=AcceptanceChecklist.Phase.FACTORY,
        )
        response = client.post(
            "/api/project-content/acceptance-criteria/bulk-update/",
            {
                "ids": [str(fat.id)],
                "changes": {"status": "Passed", "priority": "High"},
            },
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]["status"], "Passed")

        response = client.post(
            "/api/project-content/acceptance-comments/",
            {
                "criterion": str(fat.id),
                "content": "Verified during the factory test.",
                "mentions": [],
                "attachments": [],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201)
        comment_id = response.data["id"]

        response = client.post(
            f"/api/project-content/acceptance-comments/{comment_id}/reactions/toggle/",
            {"emoji": "approved"},
            format="json",
        )
        self.assertEqual(response.status_code, 201)

        response = client.get(
            "/api/project-content/audit-events/",
            {"project": self.project.id, "scope": "acceptance"},
        )
        self.assertEqual(response.status_code, 200)
        event_types = {row["event_type"] for row in response.data["results"]}
        self.assertIn("acceptance_generation_applied", event_types)
        self.assertIn("acceptance_comment_created", event_types)
